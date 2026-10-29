r"""build_specs.py — MARKET-SPECS-01 (DAY ORDER 5, 2026-10-24): one market specification (.docx) per node market, generated from the live record, plus the EU
regime spec. Every value is read at build time from the wire — nothing typed:
  · the ECO object           https://<cc>-eco.ai/eco.json (the twelve fields, each with its source page; signature and sha256)
  · the node                 the resolver's list_nodes line (state, records, events, as_of, door) and the door's facts.json (exchanges, issuers with events, region)
  · Issuerwire coverage      https://issuerwire.ai/facts.json (shown markets, issuers and signed rows per market; seeded-not-shown named apart)
  · the halts source         the example issuer's get_signal answer (halt: page, read_by, read_at) and the source research's stated limits (README "Limits")
  · ISIN coverage            the apex index (https://mcp.getissuer.ai/index.json): ISIN keys over records per node
  · the door                 the node's door, card and machine kit, by the door name rule of record
Last page: the Allooloo CEO contact page. Writes <out>\<cc>-market-spec-<date>.docx (+ eu-regime-spec) and <out>\specs-<date>.json (what was read, when).
The eco-refresh job runs it monthly after its re-sign (out = the refresh drop), so the set is regenerated with the ECO record.
Usage: python -X utf8 -I build_specs.py [out dir, default C:\ALLOOLOO\SPECS]"""
import os, sys, json, re, datetime, urllib.request, urllib.error, glob
from docx import Document
from docx.shared import Pt, RGBColor
OUT = sys.argv[1] if len(sys.argv) > 1 else r'C:\ALLOOLOO\SPECS'; os.makedirs(OUT, exist_ok=True)
HERE = os.path.dirname(os.path.abspath(__file__)); RAILS = os.path.join(HERE, '..', 'RAILS')
UA = 'Allooloo digest (MARKET-SPECS-01)'; TODAY = datetime.date.today().isoformat()
TWELVE = 'ca uk us de fr nl ch au sg jp kr hk'.split(); EU14 = 'it es se dk no fi be ie at pl pt lu gr cz'.split(); ROW = 'ky bm je in tw ae br mx'.split()   # sa out: SEEKING PARTNER (CEO ruling 2026-10-24) — 34 node markets
import sys as _s, os as _o; _s.path.insert(0, _o.path.dirname(_o.path.abspath(__file__))); import roster   # ENUM-01 (DAY ORDER 6): the roster of record
MARKETS = roster.codes()   # ENUM-01: the roster (the typed TWELVE/EU14/ROW above kept as the record only)
EU13 = [c for c in EU14 if c != 'lu']
_door = {roster.cc_of(n): n['door'].rstrip('/') for n in roster.load()}
door_of = lambda cc: _door[cc]   # the roster's door field of record
card_of = lambda cc: (_door[cc] + '/.well-known/agent-card.json') if '.eu-cm-kg.ai' in _door[cc] else _door[cc].replace('://mcp.', '://agent.') + '/.well-known/agent-card.json'
CEO = ['Matthew Keddy', 'Chief Executive Officer · Allooloo Technologies Corp.', 'mk@allooloo.ai · +1 647 237 7744 (Tel / WhatsApp)', 'allooloo.ai · issuerwire.ai · getissuer.ai']
def get(url, body=None):
    try:
        req = urllib.request.Request(url, data=json.dumps(body).encode() if body else None, headers={'User-Agent': UA, 'Accept': 'application/json', 'Content-Type': 'application/json'})
        return json.loads(urllib.request.urlopen(req, timeout=60).read().decode('utf-8'))
    except Exception as e: return {'_error': repr(e)[:160]}
def mcp(tool, args):
    j = get('https://mcp.getissuer.ai/mcp', {'jsonrpc': '2.0', 'id': 1, 'method': 'tools/call', 'params': {'name': tool, 'arguments': args}})
    try: r = j['result']; return r.get('structuredContent') or json.loads(r['content'][0]['text'])
    except Exception: return {'_error': j.get('_error') or 'no result'}
nodes = {n['node'].split('-')[0]: n for n in (mcp('list_nodes', {}).get('nodes') or [])}
index = get('https://mcp.getissuer.ai/index.json'); iw = get('https://issuerwire.ai/facts.json')
isin_by, rec_by = {}, {}
for k, v in (index.get('isin') or {}).items():
    for n in (v if isinstance(v, list) else [v]): isin_by[n] = isin_by.get(n, 0) + 1
for k, v in (index.get('cmr') or {}).items():
    for n in (v if isinstance(v, list) else [v]): rec_by[n] = rec_by.get(n, 0) + 1
EX = json.load(open(os.path.join(HERE, 'examples.json'), encoding='utf-8')) if os.path.exists(os.path.join(HERE, 'examples.json')) else {}
def limits(cc):
    for d in ('eu-sources', 'row-sources'):
        p = os.path.join(RAILS, d, cc, 'README.md')
        if os.path.exists(p):
            t = open(p, encoding='utf-8').read(); m = re.search(r'## Limits\s*\n(.*?)(\n## |\Z)', t, re.S)
            if m: return [re.sub(r'^[-*]\s*', '', l).strip() for l in m.group(1).splitlines() if l.strip().startswith(('-', '*'))][:8], f'RAILS\\{d}\\{cc}\\README.md'
    return [], None
def val(v):
    if v is None: return ''
    if isinstance(v, list): return '; '.join(val(x) for x in v)
    if isinstance(v, dict): return ' — '.join(f'{k}: {val(x)}' for k, x in v.items() if x is not None)
    return str(v)
def doc_base(title, sub):
    d = Document(); st = d.styles['Normal']; st.font.name = 'Calibri'; st.font.size = Pt(10)
    h = d.add_heading(title, 0); d.add_paragraph(sub).runs[0].font.color.rgb = RGBColor(0x55, 0x55, 0x55); return d
def table(d, rows, hdr=('Field', 'Value', 'Source')):
    t = d.add_table(rows=1, cols=len(hdr)); t.style = 'Light Grid Accent 1'
    for i, h in enumerate(hdr): t.rows[0].cells[i].text = h
    for r in rows:
        c = t.add_row().cells
        for i, x in enumerate(r): c[i].text = str(x if x is not None else '')
def contact(d):
    d.add_page_break(); d.add_heading('Contact', 1)
    for i, l in enumerate(CEO): p = d.add_paragraph(l); p.runs[0].bold = i == 0
    d.add_paragraph('Allooloo Technologies Corp. · public-record only · no prices, no quotes, no licensed market data. Artificial intelligence (AI) makes mistakes; every field names the page it was read on — read the regulator. Nothing here is investment advice.')
log = {'built': TODAY, 'out': OUT, 'markets': {}}
for cc in MARKETS:
    eco = get(f'https://{cc}-eco.ai/eco.json'); rec = (eco.get('record') or {}); f = rec.get('fields') or {}
    n = nodes.get(cc) or {}; live = bool(n.get('live')); facts = get(door_of(cc) + '/facts.json') if live else {}
    country = rec.get('country') or n.get('country') or cc
    d = doc_base(f'{country} — market specification', f'{cc}-cm-kg · generated {TODAY} from the live record (MARKET-SPECS-01) · Allooloo Technologies Corp.')
    d.add_heading('1 · Node', 1)
    table(d, [('State', 'NODE LIVE' if live else 'NODE OPENING', 'list_nodes, https://mcp.getissuer.ai/mcp'), ('As of', n.get('as_of') if live else '', 'list_nodes'),
              ('Records', n.get('records') if live else '', 'list_nodes'), ('Disclosure events', n.get('events') if live else '', 'list_nodes'),
              ('Issuers with events', facts.get('issuers_with_events'), door_of(cc) + '/facts.json' if live else ''), ('Exchanges (MIC)', val(facts.get('exchanges')), door_of(cc) + '/facts.json' if live else ''),
              ('Serving region', facts.get('region'), door_of(cc) + '/facts.json' if live else ''), ('Store', facts.get('store'), door_of(cc) + '/facts.json' if live else '')])
    d.add_heading('2 · Rules — the Electronic Compliance Object', 1)
    if f: table(d, [(k.replace('_', ' '), val(v.get('value')) if v.get('state') == 'sourced' else 'not found', (v.get('source') or v.get('url') or '') if v.get('state') == 'sourced' else ', '.join(v.get('checked') or [])) for k, v in f.items()])
    else: d.add_paragraph('No ECO record answered at build: ' + str(eco.get('_error') or 'no record'))
    sg = eco.get('signature') or {}; d.add_paragraph(f"Record sha256 {eco.get('sha256', '')} · signature {sg.get('kid', 'none')} · {f'https://{cc}-eco.ai/eco.json'}")
    d.add_heading('3 · Issuerwire coverage', 1)
    M = 'UK' if cc == 'uk' else cc.upper(); shown = M in (iw.get('markets') or []); seeded = M in (iw.get('seeded_not_shown') or [])
    table(d, [('Shown on issuerwire.ai', 'yes' if shown else ('seeded — shown after two green nights' if seeded else 'not carried'), 'https://issuerwire.ai/facts.json'),
              ('Issuers', (iw.get('issuers_by_market') or {}).get(M), 'issuerwire.ai/facts.json'), ('Signed filing rows', (iw.get('filings_by_market') or {}).get(M), 'issuerwire.ai/facts.json')])
    d.add_heading('4 · Halts — source and stated limits', 1)
    ex = (EX.get(cc) or {}).get('isin'); sig = mcp('get_signal', {'identifier': ex}) if (live and ex) else {}; h = sig.get('halt') or {}
    table(d, [('Halts page', h.get('page') or h.get('source_url') or val((f.get('trading_halts_page') or {}).get('value')), 'get_signal on ' + (ex or '—') if h else 'ECO trading_halts_page'),
              ('Read by', h.get('read_by'), 'the record’s halt flag'), ('Read at', h.get('read_at'), 'the record’s halt flag'), ('Example signal', f"{sig.get('signal', '')} · {sig.get('issuer', '')}" if sig else '', 'get_signal')])
    lim, lsrc = limits(cc)
    if lim: d.add_paragraph('Stated limits (' + lsrc + '):'); [d.add_paragraph(l, style='List Bullet') for l in lim]
    d.add_heading('5 · ISIN coverage', 1)
    nn = f'{cc}-cm-kg'; ri, ii = rec_by.get(nn, 0), isin_by.get(nn, 0)
    table(d, [('Records in the apex index', ri, 'https://mcp.getissuer.ai/index.json (cmr keys)'), ('ISIN keys', ii, 'index.json (isin keys)'), ('Coverage', f'{(100.0 * ii / ri):.1f} %' if ri else '', 'ISIN keys / records')])
    d.add_heading('6 · Door', 1)
    table(d, [('MCP door', door_of(cc) + '/mcp' if live else '— (published the day it answers)', 'door name rule of record'), ('Agent card', card_of(cc) if live else '', ''), ('Resolver', 'https://mcp.getissuer.ai/mcp', ''),
              ('Issuer page', f'https://{cc}.eu-issuer.ai/' if cc in EU13 else f'https://{cc}-issuer.ai/', ''), ('Rules page', f'https://{cc}-eco.ai/', ''), ('Receipts', 'https://mcp.rco-a2a.ai/mcp', '')], hdr=('Item', 'Address', 'Note'))
    contact(d); p = os.path.join(OUT, f'{cc}-market-spec-{TODAY}.docx'); d.save(p)
    log['markets'][cc] = {'country': country, 'live': live, 'eco': bool(f), 'issuerwire_shown': shown, 'seeded': seeded, 'isin': [ii, ri], 'halt_read': bool(h), 'file': os.path.basename(p)}
    print(cc, country, 'live' if live else 'opening', 'eco' if f else 'NO ECO', 'iw shown' if shown else ('iw seeded' if seeded else 'iw —'), f'isin {ii}/{ri}')
reg = get('https://eu-eco.ai/eco.json'); rr = reg.get('record') or {}
d = doc_base('European Union — regime specification', f'eu-eco.ai · generated {TODAY} from the live record (MARKET-SPECS-01) · Allooloo Technologies Corp.')
d.add_heading('1 · The regime', 1); table(d, [(k.replace('_', ' '), val(v.get('value')) if v.get('state') == 'sourced' else 'not found', v.get('source') or v.get('url') or '') for k, v in (rr.get('fields') or {}).items()])
d.add_heading('2 · Markets', 1); table(d, [(r.get('country'), r.get('nca'), r.get('eco_page') or '—', ('NODE LIVE · ' + str(r.get('records'))) if r.get('node_live') else 'NODE OPENING') for r in (reg.get('rows') or [])], hdr=('Market', 'Competent authority', 'ECO page', 'Node'))
contact(d); d.save(os.path.join(OUT, f'eu-regime-spec-{TODAY}.docx')); log['regime'] = bool(rr)
json.dump(log, open(os.path.join(OUT, f'specs-{TODAY}.json'), 'w', encoding='utf-8'), indent=1, ensure_ascii=False)
print('specs', len(log['markets']) + 1, 'files ->', OUT)
