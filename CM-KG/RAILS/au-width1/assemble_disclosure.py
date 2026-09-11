r"""Merges every worker's events from every pond drop of this rail (plus the last versioned event set), dedupes on (exchange, code, type, date, URL),
applies the name rule to search-found wire hits, and writes CM-KG\DISCLOSURE\au-disclosure.xlsx (Events, Coverage, Gaps, HITL, Method) and
CM-KG\DISCLOSURE\events\au-events.jsonl, with a versioned copy under POND\au-cm-kg\assembled\<date>\."""
import json, os, re, sys, datetime, shutil
from collections import Counter, defaultdict
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter
from common import *
OUT_X = sys.argv[1] if len(sys.argv) > 1 else r'C:\ALLOOLOO\CM-KG\DISCLOSURE\au-disclosure.xlsx'
OUT_J = sys.argv[2] if len(sys.argv) > 2 else r'C:\ALLOOLOO\CM-KG\DISCLOSURE\events\au-events.jsonl'
issuers = load_issuers(); byk = {key(r): r for r in issuers}
ALIASES = {}
_al = pond.latest(NODE, 'width1', 'aliases.json') or ('raw/aliases.json' if os.path.exists('raw/aliases.json') else None)
if _al: ALIASES = json.load(open(_al, encoding='utf-8'))
def load(fn):
    out = []; seen_paths = set()
    paths = [p for d, p in pond.drops(NODE, 'width1')] + (['raw'] if os.path.isdir('raw') else [])
    for p in paths:
        f = os.path.join(p, os.path.basename(fn)); rp = os.path.realpath(f)
        if not os.path.exists(f) or rp in seen_paths: continue
        seen_paths.add(rp)
        for line in open(f, encoding='utf-8'):
            try: out.append(json.loads(line))
            except Exception: pass
    return out
events = []; gaps = []; worker_notes = defaultdict(list)
PAGES = re.compile(r'\s+\d+ pages?\s+[\d.]+\s*[KM]B\s*$')
def take(rows, worker):
    for d in rows:
        k = d.get('key')
        for e in d.get('events', []):  # drops before the 2026-09-11 headline fix carried '<n> pages <size>' on ASX headlines
            if e.get('read_by', '').startswith('asx'):
                e['title'] = PAGES.sub('', e.get('title') or '').strip()
                if e.get('asx_category'): e['asx_category'] = PAGES.sub('', e['asx_category']).strip()
        if d.get('error'): gaps.append((k, worker, 'error: ' + str(d['error'])[:160])); continue
        if d.get('gap') and d['gap'] not in ('not an ASX issuer', 'not an NSX issuer'): gaps.append((k, worker, d['gap']))
        for e in d.get('events', []): events.append(e)
        if d.get('ambiguous'): worker_notes[k].append(f"{d['ambiguous']} wire hits named another issuer (not taken)")
        if d.get('note'): worker_notes[k].append(d['note'])
take(load('raw/asx_events.jsonl'), 'asx_events'); take(load('raw/asic_events.jsonl'), 'asic_events'); take(load('raw/nsx_events.jsonl'), 'nsx_events'); take(load('raw/wire_search.jsonl'), 'wire_search')
for d in load('raw/wire_search.jsonl'):
    r = byk.get(d.get('key'))
    if not r: continue
    for u in d.get('undated', []):
        dt = url_date(u['url'])
        if dt and in_window(dt): events.append(event(r, classify(u['title'], default='newswire_release'), dt, u['title'], u['wire'], u['url'], 'Tavily search (wire domains) + URL date', wire=u['wire']))
for e in load('raw/grok_live.jsonl'): events.append(e)
n_prior = 0
_prev = pond.latest_assembled(NODE, 'au-events.jsonl') or (OUT_J if os.path.exists(OUT_J) else None)
if _prev:
    for line in open(_prev, encoding='utf-8'):
        try:
            e = json.loads(line); e.pop('as_of', None); e.pop('node', None); e.pop('width', None)
            if e.get('read_by', '').startswith('asx'): e['title'] = PAGES.sub('', e.get('title') or '').strip(); e['asx_category'] = PAGES.sub('', e.get('asx_category') or '').strip()
            events.append(e); n_prior += 1
        except Exception: pass
    print('merged prior assembled events', n_prior, 'from', _prev, flush=True)
for e in events: e.pop('key', None)
events = [e for e in events if e.get('url', '').startswith('http') and e.get('date')]
kept = []; ambiguous = Counter()
for e in events:
    if 'asx.com.au' in e['read_by'] or 'ASIC' in e['read_by'] or 'nsx.com.au' in e['read_by'] or 'Grok' in e['read_by']:
        kept.append(e); continue
    k = e['exchange'] + '|' + e['ticker']
    mode, tok = name_match(e['title'], e['issuer'], tuple(ALIASES.get(k, [])))
    if mode == 'full': kept.append(e)
    elif mode == 'distinctive':
        e['detail'] = (e['detail'] + '; ' if e['detail'] else '') + f"matched on distinctive token '{tok}'"; kept.append(e)
    else:
        ambiguous[k] += 1
        why = (f"token '{tok}' is not unique across the AU roster or is a common English word" if mode == 'ambiguous' else 'no name token in the title')
        gaps.append((k, 'ambiguous match', f"{e['date']} {e['title'][:90]} | {why} | {e['url']}"))
events = kept
def prio(e): return 0 if 'asx.com.au' in e['read_by'] or 'nsx.com.au' in e['read_by'] else (1 if 'ASIC' in e['read_by'] else 2)
seen = {}
for e in sorted(events, key=prio):
    kk = (e['exchange'], e['ticker'], e['event_type'], e['date'], e['url'].lower().rstrip('/'))
    if kk not in seen: seen[kk] = e
events = sorted(seen.values(), key=lambda e: (e['exchange'], e['ticker'], e['date'], e['event_type']))
os.makedirs(os.path.dirname(OUT_J), exist_ok=True)
with open(OUT_J, 'w', encoding='utf-8') as f:
    for e in events: f.write(json.dumps({**e, 'as_of': TODAY.isoformat(), 'node': NODE, 'width': 1}, ensure_ascii=False) + '\n')
CTRL = re.compile(r'[\x00-\x08\x0b\x0c\x0e-\x1f]')
wb = Workbook(); wb.remove(wb.active)
ARIAL = Font(name='Arial', size=10); BOLD = Font(name='Arial', size=10, bold=True); HFILL = PatternFill('solid', fgColor='DDE4EE')
def sheet(name, cols, rows, widths):
    ws = wb.create_sheet(name); ws.append(cols)
    for c in ws[1]: c.font = BOLD; c.fill = HFILL; c.alignment = Alignment(wrap_text=True, vertical='top')
    for r in rows: ws.append([CTRL.sub('', v) if isinstance(v, str) else v for v in r])
    for row in ws.iter_rows(min_row=2):
        for c in row: c.font = ARIAL
    ws.freeze_panes = 'A2'; ws.auto_filter.ref = f'A1:{get_column_letter(len(cols))}{max(len(rows) + 1, 2)}'
    for i, w in enumerate(widths, 1): ws.column_dimensions[get_column_letter(i)].width = w
    return ws
ECOLS = ['Exchange', 'Code', 'ISIN', 'LEI', 'Issuer', 'Event type', 'Date', 'Title', 'ASX category', 'Price sensitive', 'Wire', 'Source', 'URL', 'Read by', 'Detail', 'State']
erows = [[e['exchange'], e['ticker'], e['isin'], e['lei'], e['issuer'], e['event_type'], e['date'], e['title'], e.get('asx_category', ''), e.get('price_sensitive', ''), e['wire'], e['source'], e['url'], e['read_by'], e['detail'], e.get('state', 'sourced')] for e in events]
sheet('Events', ECOLS, erows, [10, 8, 15, 22, 34, 20, 11, 60, 30, 10, 14, 28, 60, 40, 30, 10])
per = Counter((e['exchange'], e['ticker']) for e in events); by_type = defaultdict(Counter)
for e in events: by_type[(e['exchange'], e['ticker'])][e['event_type']] += 1
TYPES = ['asx_announcement', 'financial_statement', 'quarterly', 'agm_record_date', 'director_interest', 'substantial_holder', 'corporate_action', 'halt_resume', 'newswire_release', 'nsx_announcement', 'name_change', 'status_change']
CCOLS = ['Exchange', 'Code', 'ISIN', 'Issuer', 'ACN', 'Newswire of habit', 'Events (live)', 'Events (built)'] + TYPES + ['Price-sensitive (built)', 'Zero events (live)', 'Notes']
n_ev = len(events) + 1; crows = []
ps = Counter((e['exchange'], e['ticker']) for e in events if e.get('price_sensitive') == 'true')
for r in issuers:
    k = (r['exchange'], r['ticker']); i = len(crows) + 2
    crows.append([r['exchange'], r['ticker'], r['isin'], r['name'], r['acn'], r['wire'], f"=COUNTIFS(Events!$A$2:$A${n_ev},A{i},Events!$B$2:$B${n_ev},B{i})", per.get(k, 0)] + [by_type[k].get(t, 0) for t in TYPES] + [ps.get(k, 0), f'=IF(G{i}=0,1,0)', '; '.join(worker_notes.get(key(r), []))[:400]])
cv = sheet('Coverage', CCOLS, crows, [10, 8, 15, 34, 12, 26, 12, 12] + [12] * len(TYPES) + [12, 12, 60])
zc = get_column_letter(10 + len(TYPES))
cv.append([]); cv.append(['TOTAL', '', '', '', '', '', f'=SUM(G2:G{len(crows) + 1})', f'=SUM(H2:H{len(crows) + 1})'] + [f'=SUM({get_column_letter(9 + j)}2:{get_column_letter(9 + j)}{len(crows) + 1})' for j in range(len(TYPES) + 1)] + [f'=SUM({zc}2:{zc}{len(crows) + 1})', 'issuers with zero events (live)'])
for c in cv[cv.max_row]: c.font = BOLD
GCOLS = ['Exchange', 'Code', 'Issuer', 'Worker / field', 'Reason']; grows = []
for k, w, why in gaps:
    r = byk.get(k, {'exchange': k.split('|')[0], 'ticker': k.split('|')[1], 'name': ''}); grows.append([r['exchange'], r['ticker'], r['name'], w, why])
zero = [r for r in issuers if per.get((r['exchange'], r['ticker']), 0) == 0]
for r in zero: grows.append([r['exchange'], r['ticker'], r['name'], 'coverage', 'zero events in the window from every source that answers machines'])
BROKEN = [('TMX Australia', '', '', 'TMX Australia (formerly Cboe Australia)', 'listings and announcements sit behind reCAPTCHA with no rendered list; not attempted per ORDER-011 (HITL).'),
          ('ASX', '', '', 'ASX announcement documents', 'the viewer link resolves to the PDF behind an interstitial; the fact, date, headline and price-sensitive flag are carried, not the document.'),
          ('all', '', '', 'ASIC status-change dates', 'the bulk register carries the current status only; a non-registered status is dated on the dataset month and says so.'),
          ('NSX', '', '', 'NSX announcements feed', 'the public RSS feed holds recent announcements only; older NSX announcements need the per-company page (HTML, not attempted this run).')]
for b in BROKEN: grows.append(list(b))
gs = sheet('Gaps', GCOLS, grows, [12, 8, 34, 24, 110])
gs.append([]); gs.append(['TOTAL gap rows', '', '', '', f'=COUNTA(E2:E{len(grows) + 1})'])
HITL = [('TMX Australia', 'A machine-readable list of TMX Australia primary listings and their announcements needs the exchange (email) or a data account; the site is reCAPTCHA-gated.'),
        ('ASIC Connect API', 'Per-company register events with dates (name changes, status changes, officeholders) need an ASIC Connect / API key (AGENT KEYS\\asic.txt absent); the weekly bulk file carries current values only.'),
        ('Anthropic Admin API key', 'Balance and cost reads for the spend line need an Admin API key; the Messages key answers 401.')]
hs = sheet('HITL — needs MK', ['Item', 'What is needed'], [list(h) for h in HITL], [30, 150])
for row in hs.iter_rows(min_row=2):
    for c in row: c.alignment = Alignment(wrap_text=True, vertical='top')
METHOD = [('Order', f'ORDER-011 Australia Width 1: Disclosure. Window {SINCE} to {TODAY}. Scope: {len(issuers)} corporate issuers from au-issuers.xlsx (security type Corporate; ETFs, trusts, LICs/LITs and debt out). Pond-native: inputs are every drop of POND\\au-cm-kg\\width1 plus the last versioned event set.'),
          ('Rule', 'Blank is blank. No event without a URL. No date read out of a headline: the event date is the announcement, register or release date. Wire hits naming another issuer are counted in Coverage notes, not taken.'),
          ('Name rule', 'A search-found wire release is attributed only when its title carries every key token of the issuer name (or of a sourced alias), or one token unique across the AU roster and not a common English word; otherwise Gaps "ambiguous match". ASX and NSX exchange feeds and the ASIC register are the issuer\'s own listings and are exempt.'),
          ('ASX announcements', 'The ASX announcements platform search page per code (12-month window across two calendar years): date, headline (= the ASX category as published: Appendix 4E/4D, 4C/5B, Trading Halt, substantial holder forms, Appendix 3X/3Y/3Z, Appendix 3B, Dividend/Distribution, Notice of Meeting, …), price-sensitive marker, viewer link. Event type from the headline: financial_statement, quarterly, halt_resume, substantial_holder, director_interest, agm_record_date, corporate_action, asx_announcement.'),
          ('ASIC', 'Weekly bulk register (Width 0 drop): name changes dated from the previous-name row, deregistration dates, non-registered status (dated on the dataset month).'),
          ('NSX', 'The exchange\'s company-news RSS feed matched on the NSX code in the item title (recent announcements only).'),
          ('Newswire (search)', 'Tavily news search over PR Newswire, GlobeNewswire, Business Wire, ACCESS Newswire, Newsfile, Medianet, PRWire with published_date; a general "<short name>" announces query; date from published_date or the release URL.'),
          ('Idempotence', 'Dedupe key (exchange, code, event type, date, URL). Weekly refresh opens a new pond drop with WINDOW_DAYS=7 and re-assembles by merging every drop and the last versioned set.'),
          ('Read-by labels', 'asx.com.au announcements search (exchange site) · ASIC company register bulk dataset (data.gov.au, weekly) · nsx.com.au company news feed (exchange feed) · Tavily search (wire domains) + published_date · Tavily search (wire domains) + URL date'),
          ('State', 'sourced = one source read. Fill and Confirm (ORDER-012) add filled / confirmed / conflict.')]
mt = sheet('Method', ['Item', 'Detail'], [list(m) for m in METHOD], [30, 150])
for row in mt.iter_rows(min_row=2):
    for c in row: c.alignment = Alignment(wrap_text=True, vertical='top')
os.makedirs(os.path.dirname(OUT_X), exist_ok=True); wb.save(OUT_X)
_ad = pond.assembled(NODE); shutil.copy(OUT_X, os.path.join(_ad, 'au-disclosure.xlsx')); shutil.copy(OUT_J, os.path.join(_ad, 'au-events.jsonl'))
print('saved', OUT_X, 'and', OUT_J, '| versioned copy', _ad)
print('events', len(events)); print('by type', Counter(e['event_type'] for e in events).most_common())
print('by source', Counter(e['read_by'] for e in events).most_common())
print('price-sensitive', sum(1 for e in events if e.get('price_sensitive') == 'true'))
print('issuers with events', len(per), '/', len(issuers), '| zero events', len(zero))
print('zero-event first 20:', [f"{r['exchange']}:{r['ticker']} {r['name']}" for r in zero[:20]])
print('gap rows', len(grows), '| by worker', Counter(g[1] for g in gaps).most_common())
print('by exchange', Counter(e['exchange'] for e in events).most_common())
