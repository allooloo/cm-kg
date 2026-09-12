r"""Merges every worker's events from every pond drop of this rail (plus the last versioned event set), dedupes on (exchange, code, type, date, URL),
applies the name rule to search-found hits, and writes CM-KG\DISCLOSURE\<cc>-disclosure.xlsx (Events, Coverage, Gaps, HITL, Method) and
CM-KG\DISCLOSURE\events\<cc>-events.jsonl, with a versioned copy under POND\<node>\assembled\<date>\."""
import json, os, re, sys, datetime, shutil
from collections import Counter, defaultdict
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter
from common import *
OUT_X = sys.argv[1] if len(sys.argv) > 1 else rf'C:\ALLOOLOO\CM-KG\DISCLOSURE\{CC}-disclosure.xlsx'
OUT_J = sys.argv[2] if len(sys.argv) > 2 else rf'C:\ALLOOLOO\CM-KG\DISCLOSURE\events\{CC}-events.jsonl'
COUNTRY = {'ch': 'Switzerland', 'de': 'Germany'}[CC]; ORDER = 'ORDER-015 Part B'
REG_LABEL = {'ch': 'SHAB', 'de': 'Handelsregister'}[CC]
issuers = load_issuers(); byk = {key(r): r for r in issuers}
ALIASES = {}
_al = pond.latest(NODE, 'width1', 'aliases.json') or ('raw/aliases.json' if os.path.exists('raw/aliases.json') else None)
if _al: ALIASES = json.load(open(_al, encoding='utf-8'))
def load(fn):
    out = []; seen_paths = set()
    paths = [p for d, p in pond.drops(NODE, 'width1')] + (['raw'] if os.path.isdir('raw') else [])
    for p in paths:
        f = os.path.join(p, os.path.basename(fn)); rp = os.path.realpath(f)
        if not os.path.exists(f) or rp in seen_paths: continue
        seen_paths.add(rp)
        for line in open(f, encoding='utf-8'):
            try: out.append(json.loads(line))
            except Exception: pass
    return out
events = []; gaps = []; worker_notes = defaultdict(list)
def take(rows, worker):
    for d in rows:
        k = d.get('key')
        if d.get('error'): gaps.append((k, worker, 'error: ' + str(d['error'])[:160])); continue
        if d.get('gap'): gaps.append((k, worker, d['gap']))
        for e in d.get('events', []): events.append(e)
        if d.get('ambiguous'): worker_notes[k].append(f"{d['ambiguous']} search hits named another issuer (not taken)")
        if d.get('note'): worker_notes[k].append(d['note'])
for fn, w in (('raw/eqs_search.jsonl', 'eqs_search'), ('raw/shab_events.jsonl', 'shab_events'), ('raw/wire_search.jsonl', 'wire_search')):
    take(load(fn), w)
for d in load('raw/wire_search.jsonl'):
    r = byk.get(d.get('key'))
    if not r: continue
    for u in d.get('undated', []):
        dt = url_date(u['url'])
        if dt and in_window(dt): events.append(event(r, classify(u['title'], default='newswire_release'), dt, u['title'], u['wire'], u['url'], 'Tavily search (wire domains) + URL date', wire=u['wire']))
for e in load('raw/grok_live.jsonl'): events.append(e)
n_prior = 0
_prev = pond.latest_assembled(NODE, f'{CC}-events.jsonl') or (OUT_J if os.path.exists(OUT_J) else None)
if _prev:
    for line in open(_prev, encoding='utf-8'):
        try:
            e = json.loads(line); e.pop('as_of', None); e.pop('node', None); e.pop('width', None); events.append(e); n_prior += 1
        except Exception: pass
    print('merged prior assembled events', n_prior, 'from', _prev, flush=True)
for e in events: e.pop('key', None)
events = [e for e in events if e.get('url', '').startswith('http') and e.get('date')]
kept = []; ambiguous = Counter()
EXEMPT = ('SHAB', 'Grok')  # the register and the live layer are the issuer's own rows
for e in events:
    if any(x in e['read_by'] for x in EXEMPT):
        kept.append(e); continue
    k = e['exchange'] + '|' + e['ticker']
    r = byk.get(k, {}); mode, tok = name_match(e['title'], e['issuer'], tuple(ALIASES.get(k, [])) + ((r.get('full_name'),) if r.get('full_name') else ()))
    if mode == 'full': kept.append(e)
    elif mode == 'distinctive':
        e['detail'] = (e['detail'] + '; ' if e['detail'] else '') + f"matched on distinctive token '{tok}'"; kept.append(e)
    else:
        ambiguous[k] += 1
        why = (f"token '{tok}' is not unique across the {COUNTRY} roster or is a common word" if mode == 'ambiguous' else 'no name token in the title')
        gaps.append((k, 'ambiguous match', f"{e['date']} {e['title'][:90]} | {why} | {e['url']}"))
events = kept
def prio(e): return 0 if 'eqs-news' in e['read_by'] or 'SHAB' in e['read_by'] else 2
seen = {}
for e in sorted(events, key=prio):
    kk = (e['exchange'], e['ticker'], e['event_type'], e['date'], e['url'].lower().rstrip('/'))
    if kk not in seen: seen[kk] = e
events = sorted(seen.values(), key=lambda e: (e['exchange'], e['ticker'], e['date'], e['event_type']))
os.makedirs(os.path.dirname(OUT_J), exist_ok=True)
with open(OUT_J, 'w', encoding='utf-8') as f:
    for e in events: f.write(json.dumps({**e, 'as_of': TODAY.isoformat(), 'node': NODE, 'width': 1}, ensure_ascii=False) + '\n')
CTRL = re.compile(r'[\x00-\x08\x0b\x0c\x0e-\x1f]')
wb = Workbook(); wb.remove(wb.active)
ARIAL = Font(name='Arial', size=10); BOLD = Font(name='Arial', size=10, bold=True); HFILL = PatternFill('solid', fgColor='DDE4EE')
def sheet(name, cols, rows, widths):
    ws = wb.create_sheet(name); ws.append(cols)
    for c in ws[1]: c.font = BOLD; c.fill = HFILL; c.alignment = Alignment(wrap_text=True, vertical='top')
    for r in rows: ws.append([CTRL.sub('', v) if isinstance(v, str) else v for v in r])
    for row in ws.iter_rows(min_row=2):
        for c in row: c.font = ARIAL
    ws.freeze_panes = 'A2'; ws.auto_filter.ref = f'A1:{get_column_letter(len(cols))}{max(len(rows) + 1, 2)}'
    for i, w in enumerate(widths, 1): ws.column_dimensions[get_column_letter(i)].width = w
    return ws
ECOLS = ['Exchange', 'Code', 'ISIN', 'LEI', 'Issuer', 'Event type', 'Date', 'Title', 'Category', 'Reference', 'Language', 'Wire', 'Source', 'URL', 'Read by', 'Detail', 'State']
erows = [[e['exchange'], e['ticker'], e['isin'], e['lei'], e['issuer'], e['event_type'], e['date'], e['title'], e.get('category', ''), e.get('reference', ''), e.get('language', ''), e['wire'], e['source'], e['url'], e['read_by'], e['detail'], e.get('state', 'sourced')] for e in events]
sheet('Events', ECOLS, erows, [10, 8, 15, 22, 34, 20, 11, 60, 22, 18, 8, 14, 30, 60, 44, 30, 10])
per = Counter((e['exchange'], e['ticker']) for e in events); by_type = defaultdict(Counter)
for e in events: by_type[(e['exchange'], e['ticker'])][e['event_type']] += 1
TYPES = ['ad_hoc', 'corporate_news', 'results', 'agm_egm', 'dividend', 'halt_suspension', 'major_holder', 'directors_dealings', 'director_change', 'name_change', 'register_new_registration', 'register_mutation', 'register_deletion', 'exchange_bulletin', 'newswire_release']
CCOLS = ['Exchange', 'Code', 'ISIN', 'Issuer', 'Register id', 'Newswire of habit', 'Events (live)', 'Events (built)'] + TYPES + ['Zero events (live)', 'Notes']
n_ev = len(events) + 1; crows = []
for r in issuers:
    k = (r['exchange'], r['ticker']); i = len(crows) + 2
    crows.append([r['exchange'], r['ticker'], r['isin'], r['name'], r.get('reg_id', ''), r['wire'], f"=COUNTIFS(Events!$A$2:$A${n_ev},A{i},Events!$B$2:$B${n_ev},B{i})", per.get(k, 0)] + [by_type[k].get(t, 0) for t in TYPES] + [f'=IF(G{i}=0,1,0)', '; '.join(worker_notes.get(key(r), []))[:400]])
cv = sheet('Coverage', CCOLS, crows, [10, 8, 15, 34, 16, 26, 12, 12] + [12] * len(TYPES) + [12, 60])
zc = get_column_letter(9 + len(TYPES))
cv.append([]); cv.append(['TOTAL', '', '', '', '', '', f'=SUM(G2:G{len(crows) + 1})', f'=SUM(H2:H{len(crows) + 1})'] + [f'=SUM({get_column_letter(9 + j)}2:{get_column_letter(9 + j)}{len(crows) + 1})' for j in range(len(TYPES))] + [f'=SUM({zc}2:{zc}{len(crows) + 1})', 'issuers with zero events (live)'])
for c in cv[cv.max_row]: c.font = BOLD
GCOLS = ['Exchange', 'Code', 'Issuer', 'Worker / field', 'Reason']; grows = []
for k, w, why in gaps:
    r = byk.get(k, {'exchange': k.split('|')[0], 'ticker': k.split('|')[1], 'name': ''}); grows.append([r['exchange'], r['ticker'], r['name'], w, why])
zero = [r for r in issuers if per.get((r['exchange'], r['ticker']), 0) == 0]
for r in zero: grows.append([r['exchange'], r['ticker'], r['name'], 'coverage', 'zero events in the window from every source that answers machines'])
BROKEN = {'ch': [('all', '', '', 'SIX ad hoc publicity', 'SIX Exchange Regulation publishes ad hoc announcements through the issuers (Art. 53 LR); the SER pages on the old paths answer 404 and no public feed was found. EQS News carries most Swiss ad hoc announcements and is read here; the issuer\'s own ad hoc page is recovered in the Fill pass.'),
                 ('all', '', '', 'EQS News', 'eqs-news.com news endpoints refuse machines (404); the news pages are public and indexed, so discovery is by Tavily and coverage is bounded by its index.')],
          'de': [('all', '', '', 'EQS News', 'eqs-news.com news endpoints refuse machines (404); the news pages are public and indexed, so discovery is by Tavily and coverage is bounded by its index.'),
                 ('all', '', '', 'Handelsregister events', 'handelsregister.de is a search form without an API; register events (capital, board, name, seat changes) are not read (HITL). The SHAB-style public gazette does not exist for Germany; Bundesanzeiger is scrape-hostile.')]}[CC]
for b in BROKEN: grows.append(list(b))
gs = sheet('Gaps', GCOLS, grows, [12, 8, 34, 24, 110])
gs.append([]); gs.append(['TOTAL gap rows', '', '', '', f'=COUNTA(E2:E{len(grows) + 1})'])
HITL = {'ch': [('SIX / SER ad hoc feed', 'A machine-readable feed of SIX ad hoc announcements (SER) was not found on the public site; the exchange or SER can name the current feed.'), ('ZefixPublicREST account', 'Free account for the official Zefix API (the public web API behind zefix.ch is used meanwhile).'), ('Anthropic Admin API key', 'Balance and cost reads for the spend line need an Admin API key; the Messages key answers 401.')],
        'de': [('Handelsregister API', 'Register events and per-company sheets need a register data service (handelsregister.de is a search form).'), ('Börse Frankfurt API', 'The Regulated Market / Scale split per issuer refuses plain clients (403 CORS).'), ('Anthropic Admin API key', 'Balance and cost reads for the spend line need an Admin API key; the Messages key answers 401.')]}[CC]
hs = sheet('HITL — needs MK', ['Item', 'What is needed'], [list(h) for h in HITL], [30, 150])
for row in hs.iter_rows(min_row=2):
    for c in row: c.alignment = Alignment(wrap_text=True, vertical='top')
METHOD = [('Order', f'{ORDER} {COUNTRY} Width 1: Disclosure. Window {SINCE} to {TODAY}. Scope: {len(issuers)} corporate issuers with a home-country ISIN from {CC}-issuers.xlsx. Pond-native: inputs are every drop of POND\\{NODE}\\width1 plus the last versioned event set. Coverage is what the public sources give; nothing is padded.'),
          ('Rule', 'Blank is blank. No event without a URL. No date read out of a headline: the event date is the EQS publication date, the gazette publication date or the release date.'),
          ('Name rule', 'A search-found EQS or wire item is attributed only when its title carries every key token of the issuer name (short or GLEIF legal name, or a sourced alias), or one token unique across the roster and not a common word; otherwise Gaps "ambiguous match". Register publications are matched on the UID / register number: exempt.'),
          ('EQS News', 'Public EQS News pages found per issuer by Tavily (news search over the window + general search restricted to eqs-news.com); event type from the EQS category in the URL path (ad hoc, corporate news, voting rights, directors\' dealings, other capital market information, annual reports), date from published_date, language from the URL suffix.'),
          ({'ch': 'SHAB / FOSC', 'de': 'Register'}[CC], {'ch': 'The Swiss Official Gazette of Commerce publications API (rubric HR, keyword = UID and legal name, 12-month window): new registrations, mutations and deletions dated on publication, with the publication number.', 'de': 'Not read: handelsregister.de has no API (HITL).'}[CC]),
          ('Newswire (search)', 'Tavily news search over PR Newswire, GlobeNewswire, Business Wire, ACCESS Newswire, Newsfile, pressetext, IRW-Press with published_date; a general "<short name>" announces query; date from published_date or the release URL.'),
          ('Languages', 'German, French and Italian items are kept and labelled by language; the Fill pass reads them with Mistral ("read by Mistral (de/fr)").'),
          ('Idempotence', 'Dedupe key (exchange, code, event type, date, URL). Weekly refresh opens a new pond drop with WINDOW_DAYS=7 and re-assembles by merging every drop and the last versioned set.'),
          ('State', 'sourced = one source read. Fill and Confirm add filled / confirmed / conflict.')]
mt = sheet('Method', ['Item', 'Detail'], [list(m) for m in METHOD], [30, 150])
for row in mt.iter_rows(min_row=2):
    for c in row: c.alignment = Alignment(wrap_text=True, vertical='top')
os.makedirs(os.path.dirname(OUT_X), exist_ok=True); wb.save(OUT_X)
_ad = pond.assembled(NODE); shutil.copy(OUT_X, os.path.join(_ad, f'{CC}-disclosure.xlsx')); shutil.copy(OUT_J, os.path.join(_ad, f'{CC}-events.jsonl'))
print('saved', OUT_X, 'and', OUT_J, '| versioned copy', _ad)
print('events', len(events)); print('by type', Counter(e['event_type'] for e in events).most_common())
print('by source', Counter(e['read_by'] for e in events).most_common()); print('by language', Counter(e.get('language') or '' for e in events).most_common(5))
print('issuers with events', len(per), '/', len(issuers), '| zero events', len(zero))
print('zero-event first 20:', [f"{r['exchange']}:{r['ticker']} {r['name']}" for r in zero[:20]])
print('gap rows', len(grows), '| by worker', Counter(g[1] for g in gaps).most_common())
