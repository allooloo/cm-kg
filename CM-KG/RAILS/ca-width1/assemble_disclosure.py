r"""Merges every worker's events, dedupes on (ticker, type, date, URL), re-classifies wire releases from their titles,
and writes CM-KG\DISCLOSURE\ca-disclosure.xlsx (Events, Coverage, Gaps, Method) and CM-KG\DISCLOSURE\events\ca-events.jsonl."""
import json, os, sys, datetime
from collections import Counter, defaultdict
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter
from common import *
OUT_X = sys.argv[1] if len(sys.argv) > 1 else r'C:\ALLOOLOO\CM-KG\DISCLOSURE\ca-disclosure.xlsx'
OUT_J = sys.argv[2] if len(sys.argv) > 2 else r'C:\ALLOOLOO\CM-KG\DISCLOSURE\events\ca-events.jsonl'
issuers = load_issuers(); byk = {key(r): r for r in issuers}
import pond
NODE = 'ca-cm-kg'
def load(fn):
    """pond rule 3: every drop of this rail (POND\\ca-cm-kg\\width1\\<date>\\) plus the working folder; the dedupe key keeps one copy per event"""
    out = []; seen_paths = set()
    paths = [p for d, p in pond.drops(NODE, 'width1')] + (['raw'] if os.path.isdir('raw') else [])
    for p in paths:
        f = os.path.join(p, os.path.basename(fn)); rp = os.path.realpath(f)
        if not os.path.exists(f) or rp in seen_paths: continue
        seen_paths.add(rp)
        for line in open(f, encoding='utf-8'):
            try: out.append(json.loads(line))
            except Exception: pass
    return out
events = []; gaps = []; src_stats = Counter(); worker_notes = defaultdict(list)
def take(rows, worker):
    for d in rows:
        k = d.get('key')
        if d.get('error'): gaps.append((k, worker, 'error: ' + str(d['error'])[:160])); continue
        if d.get('gap') and d['gap'] != 'skipped in this pass': gaps.append((k, worker, d['gap']))
        for e in d.get('events', []): events.append(e)
        if d.get('ambiguous'): worker_notes[k].append(f"{d['ambiguous']} wire hits named another issuer (not taken)")
        if d.get('note'): worker_notes[k].append(d['note'])
take(load('raw/wire_pages.jsonl'), 'wire_pages'); take(load('raw/wire_search.jsonl'), 'wire_search'); take(load('raw/release_dates.jsonl'), 'release_dates'); take(load('raw/tsxv_bulletins.jsonl'), 'tsxv_bulletins'); take(load('raw/newsfile_snippets.jsonl'), 'newsfile_snippets')
# search hits that carried no published_date but whose URL carries one (GlobeNewswire single-digit months, Business Wire)
for d in load('raw/wire_search.jsonl'):
    r = byk.get(d.get('key'))
    if not r: continue
    for u in d.get('undated', []):
        dt = url_date(u['url'])
        if dt and in_window(dt) and not re.search(r'/Resource/Download|attachmentng|/Tracker', u['url']):
            events.append(event(r, classify(u['title']), dt, u['title'], u['wire'], u['url'], 'Tavily search (wire domains) + URL date', wire=u['wire']))
for e in load('raw/cse_bulletins.jsonl'): events.append(e)
for e in load('raw/grok_live.jsonl'): events.append(e)   # 48-hour live layer (read by Grok (live)), present on daily runs only
# pond rule 3: the last versioned event set is merged too (history is never dropped)
_prev = pond.latest_assembled(NODE, 'ca-events.jsonl') or (OUT_J if os.path.exists(OUT_J) else None)
if _prev:
    _n = 0
    for line in open(_prev, encoding='utf-8'):
        try:
            e = json.loads(line); e.pop('as_of', None); e.pop('node', None); e.pop('width', None); events.append(e); _n += 1
        except Exception: pass
    print('merged prior assembled events', _n, 'from', _prev, flush=True)
# re-classify wire releases from titles (rules live in common.py), keep bulletin types as set by their workers
for e in events:
    if e['read_by'].startswith(('Tavily', 'newsfilecorp', 'newswire.ca', 'prnewswire')): e['event_type'] = classify(e['title'])
    e['read_by'] = e['read_by'].replace('+ URL date (URL date)', '+ URL date').replace('(wire domains) (URL date)', '(wire domains) + URL date')
    e.pop('key', None)
# no event without a URL or a date
events = [e for e in events if e.get('url', '').startswith('http') and e.get('date')]
# name rule (CEO, 2026-09-10): a wire hit is the issuer's when the title carries every key token of the name, or one token that is
# unique across the full 4,820-row roster AND not a common English/French word. Anything else is an ambiguous match -> Gaps.
# Company-page reads are exempt: the page is the issuer's own listing.
kept = []; ambiguous = Counter()
for e in events:
    if 'company page' in e['read_by'] or e['read_by'] in ('apps.tmx.com TSXV company documents', 'thecse.com bulletins API'):
        kept.append(e); continue
    mode, tok = name_match(e['title'], e['issuer'])
    if mode == 'full': kept.append(e)
    elif mode == 'distinctive':
        e['detail'] = (e['detail'] + '; ' if e['detail'] else '') + f"matched on distinctive token '{tok}'"; kept.append(e)
    else:
        k = e['exchange'] + '|' + e['ticker']; ambiguous[k] += 1
        why = (f"token '{tok}' is not unique across the roster or is a common English/French word" if mode == 'ambiguous' else 'no name token in the title')
        gaps.append((k, 'ambiguous match', f"{e['date']} {e['title'][:90]} | {why} | {e['url']}"))
events = kept
# dedupe on (ticker, type, date, URL); prefer the wire company page read over a search read
def prio(e): return 0 if 'company page' in e['read_by'] else (1 if 'bulletins' in e['read_by'] or 'company documents' in e['read_by'] else 2)
seen = {}
for e in sorted(events, key=prio):
    kk = (e['exchange'], e['ticker'], e['event_type'], e['date'], e['url'].lower().rstrip('/'))
    if kk not in seen: seen[kk] = e
events = sorted(seen.values(), key=lambda e: (e['exchange'], e['ticker'], e['date'], e['event_type']))
# jsonl
os.makedirs(os.path.dirname(OUT_J), exist_ok=True)
with open(OUT_J, 'w', encoding='utf-8') as f:
    for e in events: f.write(json.dumps({**e, 'as_of': TODAY.isoformat(), 'node': 'ca-cm-kg', 'width': 1}, ensure_ascii=False) + '\n')
# workbook
wb = Workbook(); wb.remove(wb.active)
ARIAL = Font(name='Arial', size=10); BOLD = Font(name='Arial', size=10, bold=True); HFILL = PatternFill('solid', fgColor='DDE4EE')
def sheet(name, cols, rows, widths):
    ws = wb.create_sheet(name); ws.append(cols)
    for c in ws[1]: c.font = BOLD; c.fill = HFILL; c.alignment = Alignment(wrap_text=True, vertical='top')
    for r in rows: ws.append(r)
    for row in ws.iter_rows(min_row=2):
        for c in row: c.font = ARIAL
    ws.freeze_panes = 'A2'; ws.auto_filter.ref = f'A1:{get_column_letter(len(cols))}{max(len(rows) + 1, 2)}'
    for i, w in enumerate(widths, 1): ws.column_dimensions[get_column_letter(i)].width = w
    return ws
ECOLS = ['Exchange', 'Ticker', 'ISIN', 'LEI', 'Issuer', 'Event type', 'Date', 'Title', 'Wire', 'Source', 'URL', 'Read by', 'Detail']
erows = [[e['exchange'], e['ticker'], e['isin'], e['lei'], e['issuer'], e['event_type'], e['date'], e['title'], e['wire'], e['source'], e['url'], e['read_by'], e['detail']] for e in events]
sheet('Events', ECOLS, erows, [12, 10, 15, 22, 34, 18, 11, 60, 18, 24, 60, 40, 30])
# Coverage: per issuer, events by type, with live COUNTIFS over Events
per = Counter((e['exchange'], e['ticker']) for e in events)
by_type = defaultdict(Counter)
for e in events: by_type[(e['exchange'], e['ticker'])][e['event_type']] += 1
TYPES = ['newswire_release', 'financial_statement', 'agm_record_date', 'early_warning', 'corporate_action', 'halt_resume', 'exchange_bulletin']
CCOLS = ['Exchange', 'Ticker', 'ISIN', 'Issuer', 'Newswire of habit', 'Events (live)', 'Events (built)'] + TYPES + ['Zero events (live)', 'Notes']
n_ev = len(events) + 1
crows = []
for r in issuers:
    k = (r['exchange'], r['ticker']); i = len(crows) + 2
    crows.append([r['exchange'], r['ticker'], r['isin'], r['name'], r['wire'], f"=COUNTIFS(Events!$A$2:$A${n_ev},A{i},Events!$B$2:$B${n_ev},B{i})", per.get(k, 0)] + [by_type[k].get(t, 0) for t in TYPES] + [f'=IF(F{i}=0,1,0)', '; '.join(worker_notes.get(key(r), []))[:400]])
cv = sheet('Coverage', CCOLS, crows, [12, 10, 15, 34, 20, 12, 12] + [14] * len(TYPES) + [12, 60])
cv.append([]); cv.append(['TOTAL', '', '', '', '', f'=SUM(F2:F{len(crows) + 1})', f'=SUM(G2:G{len(crows) + 1})'] + [f'=SUM({get_column_letter(8 + j)}2:{get_column_letter(8 + j)}{len(crows) + 1})' for j in range(len(TYPES))] + [f'=SUM({get_column_letter(8 + len(TYPES))}2:{get_column_letter(8 + len(TYPES))}{len(crows) + 1})', 'issuers with zero events (live COUNTIF): ' + f'=COUNTIF({get_column_letter(8 + len(TYPES))}2:{get_column_letter(8 + len(TYPES))}{len(crows) + 1},1)'])
for c in cv[cv.max_row]: c.font = BOLD
# Gaps
GCOLS = ['Exchange', 'Ticker', 'Issuer', 'Worker / field', 'Reason']
grows = []
for k, w, why in gaps:
    r = byk.get(k, {'exchange': k.split('|')[0], 'ticker': k.split('|')[1], 'name': ''}); grows.append([r['exchange'], r['ticker'], r['name'], w, why])
zero = [r for r in issuers if per.get((r['exchange'], r['ticker']), 0) == 0]
for r in zero: grows.append([r['exchange'], r['ticker'], r['name'], 'coverage', 'zero events in the window from every source that answers machines'])
BROKEN = [('all', '', '', 'SEDI / early-warning system', 'not attempted per ORDER-003: sedi.ca answers HTTP 403 to machines; early-warning press releases on the wires are carried as events where they exist'),
          ('all', '', '', 'CIRO halts', 'ciro.ca/newsroom/halts-and-resumptions answers HTTP 403 to machines and the browser navigation is denied; halts and resumes come from TSXV and CSE bulletins and wire releases only'),
          ('all', '', '', 'CDS bulletins', 'cds.ca bulletin lists sit behind the participant login; corporate-action dates come from exchange bulletins and wire releases only'),
          ('TSX', '', '', 'TSX bulletins', 'no public per-issuer bulletin list found for TSX (senior) issuers; the TMX company-documents controller serves TSXV only'),
          ('Cboe Canada', '', '', 'Cboe Canada notices', 'cboe.com listing-notices pages render no list to machines and corporate-action bulletins are a subscription service'),
          ('all', '', '', 'Newsfile company page', 'lists the latest 20 releases only; older Newsfile releases come from search and release-page dates'),
          ('all', '', '', 'GlobeNewswire / Business Wire', 'block direct fetch; releases come from search results, dated from the release URL')]
for b in BROKEN: grows.append(list(b))
gs = sheet('Gaps', GCOLS, grows, [12, 10, 34, 24, 110])
gs.append([]); gs.append(['TOTAL gap rows', '', '', '', f'=COUNTA(E2:E{len(grows) + 1})'])
# Method
MCOLS = ['Item', 'Detail']
METHOD = [('Order', f'ORDER-003 Canada Width 1: Disclosure. Window {SINCE} to {TODAY}. Scope: {len(issuers)} corporate issuers from ca-issuers.xlsx.'),
          ('Rule', 'Blank is blank. No event without a URL. No date read out of a headline: the event date is the release or bulletin date. Wire hits naming another issuer are counted in Coverage notes, not taken.'),
          ('Name rule (CEO, 2026-09-10)', 'A wire release is attributed to an issuer only when its title carries every key token of the issuer name, or a single token that (a) is unique across the full 4,820-row roster of ca-issuers.xlsx and (b) is not a common English or French word (google-10000-english plus the top 15,000 French frequency words, kept in the rail as common_words.txt). A title matched only on a token that fails either test is written to Gaps as "ambiguous match", never to Events. Events taken on a distinctive token say so in the Detail column. Reads from an issuer\'s own wire company page are exempt.'),
          ('Newswire (company pages)', 'Newsfile, Cision (newswire.ca) and PR Newswire per-company listing pages, dated by the wire; Cision and PR Newswire paged back to the window start; Newsfile shows the latest 20.'),
          ('Newswire (search)', 'Tavily news search over the seven wire domains with published_date; GlobeNewswire and Business Wire general search with the date from the release URL; undated hits dated from the release page.'),
          ('Exchange bulletins', 'TSXV: TMX company-documents controller per PO id (bulletin category, type, date, notice). CSE: exchange bulletins feed filtered to the issuer symbols. TSX and Cboe: none available (see Gaps).'),
          ('Event types', 'newswire_release · financial_statement · agm_record_date · early_warning · corporate_action · halt_resume · exchange_bulletin. Wire releases are typed from their titles with the rules in common.py; bulletins from their category and type.'),
          ('Idempotence', 'Dedupe key (exchange, ticker, event type, date, URL). Daily refresh runs the same workers with WINDOW_DAYS=2 and merges on that key.'),
          ('Read-by labels', 'newsfilecorp.com company page · newswire.ca company page · prnewswire.com company page · Tavily search (wire domains) + published_date · Tavily search (wire domains) + URL date · Tavily search (wire domains) + release page date · apps.tmx.com TSXV company documents · thecse.com bulletins API')]
mt = sheet('Method', MCOLS, [list(m) for m in METHOD], [30, 150])
for row in mt.iter_rows(min_row=2):
    for c in row: c.alignment = Alignment(wrap_text=True, vertical='top')
os.makedirs(os.path.dirname(OUT_X), exist_ok=True); wb.save(OUT_X)
import pond, shutil as _sh; _ad = pond.assembled('ca-cm-kg'); _sh.copy(OUT_X, os.path.join(_ad, 'ca-disclosure.xlsx')); _sh.copy(OUT_J, os.path.join(_ad, 'ca-events.jsonl')); print('versioned copy', _ad)  # pond rule 2
# report
print('saved', OUT_X, 'and', OUT_J)
print('events', len(events)); print('by type', Counter(e['event_type'] for e in events).most_common())
print('by source', Counter(e['read_by'] for e in events).most_common())
print('issuers with events', len(per), '/', len(issuers), '| zero events', len(zero))
print('zero-event first 20:', [f"{r['exchange']}:{r['ticker']} {r['name']}" for r in zero[:20]])
print('gap rows', len(grows), '| by worker', Counter(g[1] for g in gaps).most_common())
print('by exchange', Counter(e['exchange'] for e in events).most_common())
