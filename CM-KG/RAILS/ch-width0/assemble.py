"""Step 7 — the workbook: Coverage (live COUNTA), Gaps, SIX, BX Swiss, Gaps detail, HITL, Method. Pond-native: every input is read from the pond drops of node
ch-cm-kg (newest drop wins per key); the output is written to POND\\ch-cm-kg\\assembled\\<date>\\ch-issuers.xlsx and mirrored to CM-KG\\ISSUERS\\ch-issuers.xlsx.
Every enriched cell carries a source URL, a read-by label and a State. Nothing is inferred."""
import json, os, re, sys, datetime, shutil
from collections import Counter
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter
import pond
NODE = 'ch-cm-kg'; TODAY = datetime.date.today().isoformat()
MIRROR = sys.argv[1] if len(sys.argv) > 1 else r'C:\ALLOOLOO\CM-KG\ISSUERS\ch-issuers.xlsx'
def key(r): return r['exchange'] + '|' + r['symbol']
rows = json.load(open(pond.latest(NODE, 'width0', 'roster.json'), encoding='utf-8'))
def by_key(source, fn, k='key'): return {d[k]: d for d in pond.read_jsonl_all(NODE, source, fn, key=k)}
E_REG = by_key('width0', 'enr_reg.jsonl'); E_WIRE = by_key('width0', 'enr_wire.jsonl'); ZEFIX = by_key('width0', 'zefix.jsonl')
E_LEIREC = by_key('width0', 'lei_records.jsonl'); E_MATCH = by_key('width0', 'lei_match.jsonl', 'name')
ISIN_LEI = pond.read_json_latest(NODE, 'width0', 'isin_lei_hits.json') or {}
ISIN_META = pond.read_json_latest(NODE, 'width0', 'isin_lei_meta.json') or {}
ISIN_LEI_SRC = 'https://mapping.gleif.org/api/v2/isin-lei/latest (' + ISIN_META.get('file', 'daily file') + ')'
ZEFIX_RA = ('RA000548', 'RA000549')
def name_key(s):
    s = (s or '').upper().replace('&', ' AND ').replace('.', '').replace(',', '').replace('-', ' ')
    s = re.sub(r'\b(AG|SA|LTD|LIMITED|INC|PLC|NV|N|I|HOLDING|HOLDINGS|GROUP|GROUPE|GRUPPE|SOCIETE ANONYME|AKTIENGESELLSCHAFT|CORP|CORPORATION|CO|REG|BR|PC|PS)\b', ' ', s)
    return ' '.join(re.sub(r'[^A-Z0-9 ]', ' ', s).split())
def is_corp(r): return r['security_type'] == 'Corporate'
RB = {'six': 'SIX Swiss Exchange share explorer feed (exchange list)', 'bx': 'BX Swiss instruments list (exchange list)', 'gleif_map': 'GLEIF ISIN-to-LEI mapping file (exact ISIN match)', 'gleif_exact': 'GLEIF name-exact', 'gleif_rec': 'GLEIF LEI record (LEI per the LEI column)',
      'zefix': 'Zefix federal commercial register index (zefix.ch public web API, name-exact + corporation legal form)', 'tavily': 'Tavily search + regex extraction from page text', 'tavily_wire': 'Tavily search restricted to wire domains (release hits)'}
COLS = ['Legal name', 'Symbol', 'Valor', 'Exchange', 'Security type', 'Share type', 'ISIN', 'ISIN source', 'ISIN read by', 'ISIN State',
        'LEI', 'LEI source', 'LEI read by', 'LEI State', 'LEI registration status',
        'UID (Zefix)', 'CH-ID', 'Zefix link (cantonal excerpt)', 'Zefix source', 'Zefix read by', 'Zefix State', 'Legal form (Zefix)', 'Zefix status', 'Last SHAB date',
        'Legal seat', 'Legal seat source', 'Legal seat read by', 'Legal seat State',
        'Registered office', 'Registered office source', 'Registered office read by', 'Registered office State',
        'Incorporation jurisdiction', 'Jurisdiction source', 'Jurisdiction read by', 'Jurisdiction State',
        'Sector (SIX ICB)', 'Sector source', 'Sector read by', 'Sector State',
        'Auditor', 'Auditor source', 'Auditor read by', 'Auditor State', 'Annual report', 'Annual report source', 'Annual report read by',
        'Share registrar', 'Share registrar source', 'Share registrar read by', 'Share registrar State', 'Share registrar evidence',
        'Newswire of habit', 'Newswire releases seen', 'Newswire read by', 'Newswire State',
        'HQ city', 'HQ source', 'HQ read by', 'HQ State', 'Listing date', 'Listing date source', 'Trading currency', 'Number in issue', 'Status', 'Roster source', 'Roster read by', 'Gaps']
def build(r):
    k = key(r); ex = r['exchange']; corp = is_corp(r); gaps = []; o = {c: '' for c in COLS}; swiss = r['isin'].startswith('CH')
    o['Legal name'] = r['name']; o['Symbol'] = r['symbol']; o['Valor'] = r.get('valor') or ''; o['Exchange'] = ex; o['Security type'] = r['security_type'] + ('' if swiss or not corp else ' (foreign ISIN, SIX line)'); o['Share type'] = r.get('share_type') or ''
    o['Status'] = r['status']; o['Roster source'] = r.get('bx_url') or r['roster_src']; o['Roster read by'] = RB['six'] if ex == 'SIX' else RB['bx']; o['Trading currency'] = r.get('currency') or ''; o['Number in issue'] = r.get('number_in_issue') or ''
    if r.get('listing_date'): o['Listing date'] = r['listing_date']; o['Listing date source'] = r['roster_src']
    else: gaps.append(('Listing date', 'not on the exchange list'))
    if r.get('isin'): o['ISIN'] = r['isin']; o['ISIN source'] = o['Roster source']; o['ISIN read by'] = o['Roster read by']; o['ISIN State'] = 'sourced'
    else: gaps.append(('ISIN', 'not on the exchange list'))
    lei = ''; m = E_MATCH.get(r['name']); exact = [(l, v) for l, v in (m or {}).get('matches', [])]
    if r.get('isin') and ISIN_LEI.get(r['isin']):
        lei = ISIN_LEI[r['isin']].split('|')[0]; o['LEI'] = lei; o['LEI source'] = ISIN_LEI_SRC; o['LEI read by'] = RB['gleif_map']; o['LEI State'] = 'sourced'
        mrec = E_LEIREC.get(lei)
        if mrec and name_key(mrec['name']) != name_key(r['name']) and not (name_key(mrec['name']).startswith(name_key(r['name'])) or name_key(r['name']).startswith(name_key(mrec['name']))): o['LEI State'] = 'conflict'; gaps.append(('LEI', f"mapping file maps ISIN {r['isin']} to LEI {lei}, whose GLEIF legal name is \"{mrec['name']}\", not the issuer; LEI kept with State conflict"))
    elif exact:
        lei = exact[0][0]; o['LEI'] = lei; o['LEI source'] = f'https://api.gleif.org/api/v1/lei-records/{lei}'; o['LEI read by'] = RB['gleif_exact']; o['LEI State'] = 'sourced'
    else: gaps.append(('LEI', ('ISIN not in the GLEIF ISIN-to-LEI mapping file; ' if r.get('isin') else 'no ISIN; ') + ('no GLEIF legal name equals the roster name' if m else 'GLEIF name lookup not run')))
    rec = E_LEIREC.get(lei, {}) if lei else {}
    if lei: o['LEI registration status'] = rec.get('reg_status', '') if rec else 'record not read'
    # Zefix
    z = ZEFIX.get(k) or {}
    if z.get('uid'):
        o['UID (Zefix)'] = z['uid']; o['CH-ID'] = z.get('chid') or ''; o['Zefix link (cantonal excerpt)'] = z.get('excerpt_url') or ''; o['Zefix source'] = z.get('src', ''); o['Zefix read by'] = RB['zefix']; o['Zefix State'] = 'sourced'
        o['Legal form (Zefix)'] = z.get('legal_form') or ''; o['Zefix status'] = z.get('status') or ''; o['Last SHAB date'] = z.get('shab_date') or ''
        if z.get('legal_seat'): o['Legal seat'] = z['legal_seat']; o['Legal seat source'] = z.get('src', ''); o['Legal seat read by'] = RB['zefix']; o['Legal seat State'] = 'sourced'
        o['Incorporation jurisdiction'] = 'Switzerland (Zefix-registered ' + (z.get('legal_form') or 'corporation').lower() + ')'; o['Jurisdiction source'] = z.get('src', ''); o['Jurisdiction read by'] = RB['zefix'] + ' (legal form)'; o['Jurisdiction State'] = 'sourced'
        ras = re.sub(r'[^A-Z0-9]', '', (rec.get('registeredAs') or '').upper()); zu = re.sub(r'[^A-Z0-9]', '', z['uid'].upper())
        if rec.get('registeredAt') in ZEFIX_RA and ras and ras != zu: o['Zefix State'] = 'conflict'; gaps.append(('UID (Zefix)', f"Zefix name-exact UID {z['uid']} but the LEI record registeredAs {rec.get('registeredAs')}"))
    else:
        if swiss and corp: gaps.append(('UID (Zefix)', ('no corporation in the Zefix index equals the roster name (' + str(z.get('n_cands', 0)) + ' candidates; near: ' + '; '.join(x[0] for x in z.get('near', [])[:2]) + ')') if z and not z.get('error') else ('Zefix search error ' + str(z.get('error')) if z else ('not a Swiss ISIN' if not swiss else 'not searched'))))
        elif not swiss: gaps.append(('UID (Zefix)', 'foreign ISIN: not a Swiss register entry'))
        if rec.get('jur'): o['Incorporation jurisdiction'] = rec['jur']; o['Jurisdiction source'] = rec['src']; o['Jurisdiction read by'] = RB['gleif_rec']; o['Jurisdiction State'] = 'sourced'
        else: gaps.append(('Incorporation jurisdiction', 'no Zefix entry and no LEI record'))
    if o['Incorporation jurisdiction'].startswith('Switzerland') and rec.get('jur') and not rec['jur'].startswith('CH'): o['Jurisdiction State'] = 'conflict'; gaps.append(('Incorporation jurisdiction', f"Zefix corporation but GLEIF legal jurisdiction {rec['jur']}"))
    if rec.get('legal_city') or rec.get('legal_lines'):
        o['Registered office'] = ', '.join(v for v in ((rec.get('legal_lines') or []) + [rec.get('legal_city'), rec.get('legal_region'), rec.get('legal_postal'), rec.get('legal_country')]) if v); o['Registered office source'] = rec['src']; o['Registered office read by'] = RB['gleif_rec'] + ' (legal address)'; o['Registered office State'] = 'sourced'
    else: gaps.append(('Registered office', 'no LEI record (Zefix carries the legal seat, not the street address, in the index)'))
    if rec.get('hq_city'): o['HQ city'] = rec['hq_city']; o['HQ source'] = rec['src']; o['HQ read by'] = RB['gleif_rec'] + ' (headquarters address)'; o['HQ State'] = 'sourced'
    elif z.get('legal_seat'): o['HQ city'] = z['legal_seat']; o['HQ source'] = z.get('src', ''); o['HQ read by'] = RB['zefix'] + ' (legal seat)'; o['HQ State'] = 'sourced'
    else: gaps.append(('HQ city', 'no LEI record and no Zefix entry'))
    if r.get('sector'): o['Sector (SIX ICB)'] = r['sector']; o['Sector source'] = r['roster_src']; o['Sector read by'] = RB['six']; o['Sector State'] = 'sourced'
    else: gaps.append(('Sector (SIX ICB)', 'the SIX feed returns an empty ICB industry for this line'))
    if corp: gaps.append(('Auditor', 'Width 0 records the annual report link only (ORDER-015: source link, not a read); no machine-readable annual-report index for SIX issuers — the Fill pass reads the issuer site')); gaps.append(('Annual report', 'no exchange-hosted annual-report index (SIX does not host issuer reports); Fill pass'))
    c = E_REG.get(k)
    if c and c.get('reg'): o['Share registrar'] = c['reg']; o['Share registrar source'] = c['reg_src']; o['Share registrar read by'] = RB['tavily']; o['Share registrar evidence'] = c.get('reg_ev', ''); o['Share registrar State'] = 'sourced'
    else: gaps.append(('Share registrar', (c.get('reg_gap') or 'none found') if c else ('not searched (foreign ISIN / non-corporate)' if not (corp and swiss) else 'not searched')))
    w = E_WIRE.get(k)
    if w and w.get('wire'): o['Newswire of habit'] = w['wire']; o['Newswire releases seen'] = '\n'.join(h['url'] for h in w['hits']); o['Newswire read by'] = RB['tavily_wire']; o['Newswire State'] = 'sourced'
    else: gaps.append(('Newswire of habit', ((w.get('gap') or 'search error') if w else ('not searched (foreign ISIN / non-corporate)' if not (corp and swiss) else 'not searched')) + '; SIX ad hoc publicity is read at Width 1'))
    o['Gaps'] = '; '.join(f'{f}: {why}' for f, why in gaps)
    return o, gaps
CTRL = re.compile(r'[\x00-\x08\x0b\x0c\x0e-\x1f]')
wb = Workbook(); wb.remove(wb.active)
ARIAL = Font(name='Arial', size=10); BOLD = Font(name='Arial', size=10, bold=True); HFILL = PatternFill('solid', fgColor='DDE4EE')
SHEETS = ['SIX', 'BX Swiss']; sheets = {s: [] for s in SHEETS}; gap_detail = []; gap_summary = Counter()
for r in rows:
    o, gaps = build(r); sheets[r['exchange']].append(o)
    for f, why in gaps: gap_detail.append((r['exchange'], r['symbol'], r['name'], r['security_type'], f, why)); gap_summary[(r['exchange'], f, why[:120])] += 1
sheet_rows = {}
for ex, data in sheets.items():
    ws = wb.create_sheet(ex); ws.append(COLS)
    for cell in ws[1]: cell.font = BOLD; cell.fill = HFILL; cell.alignment = Alignment(wrap_text=True, vertical='top')
    data.sort(key=lambda o: (0 if o['ISIN'].startswith('CH') else 1, o['Legal name'].lower()))
    for o in data: ws.append([CTRL.sub('', o[c]) if isinstance(o[c], str) else o[c] for c in COLS])
    for row in ws.iter_rows(min_row=2):
        for cell in row: cell.font = ARIAL
    ws.freeze_panes = 'A2'; ws.auto_filter.ref = f"A1:{get_column_letter(len(COLS))}{max(len(data) + 1, 2)}"
    for i, c in enumerate(COLS, 1):
        w = 14
        if c in ('Legal name', 'Share registrar', 'Registered office', 'Incorporation jurisdiction', 'Newswire of habit', 'Legal form (Zefix)'): w = 34
        if 'source' in c.lower() or 'link' in c.lower() or 'releases' in c.lower(): w = 44
        if 'read by' in c.lower(): w = 30
        if 'evidence' in c.lower() or c == 'Gaps': w = 60
        if c.endswith('State'): w = 10
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[1].height = 30; sheet_rows[ex] = len(data)
cv = wb.create_sheet('Coverage', 0)
cv.append(['Field', 'Column'] + [f'{ex} filled' for ex in SHEETS] + [f'{ex} fill %' for ex in SHEETS] + ['All filled', 'All fill %'])
for cell in cv[1]: cell.font = BOLD; cell.fill = HFILL
cv.append(['Entities (rows)', 'A'] + [f"=COUNTA('{ex}'!A2:A{sheet_rows[ex] + 1})" for ex in SHEETS] + ['', ''] + ["=SUM(C2:D2)", ''])
FIELDS = ['Legal name', 'Symbol', 'Security type', 'ISIN', 'LEI', 'UID (Zefix)', 'Zefix link (cantonal excerpt)', 'Legal seat', 'Registered office', 'Incorporation jurisdiction', 'Sector (SIX ICB)', 'Auditor', 'Annual report', 'Share registrar', 'Newswire of habit', 'HQ city', 'Listing date']
rix = 3
for f in FIELDS:
    col = get_column_letter(COLS.index(f) + 1); line = [f, col]
    for ex in SHEETS: line.append(f"=COUNTA('{ex}'!{col}2:{col}{max(sheet_rows[ex] + 1, 2)})")
    for j, ex in enumerate(SHEETS): line.append(f"=IF({get_column_letter(3 + j)}$2=0,0,{get_column_letter(3 + j)}{rix}/{get_column_letter(3 + j)}$2)")
    line.append(f"=SUM(C{rix}:D{rix})"); line.append(f"=IF(G$2=0,0,G{rix}/G$2)"); cv.append(line); rix += 1
for row in cv.iter_rows(min_row=2):
    for cell in row:
        cell.font = ARIAL
        if cell.column in (5, 6, 8): cell.number_format = '0.0%'
cv.append([]); cv.append(['Note', 'Counts are live COUNTA formulas over the exchange tabs. A cell counts as filled only when a source was read; blanks are blanks (nothing inferred). Every enriched field carries source, read-by and State. Sweep date ' + TODAY + '. Inputs: the pond drops of node ch-cm-kg. SIX lines with a foreign ISIN are carried as exchange lines outside the Swiss corporate scope.'])
for cell in cv[cv.max_row]: cell.font = ARIAL
cv.freeze_panes = 'A2'
gs = wb.create_sheet('Gaps', 1); gs.append(['Exchange', 'Field', 'Reason', 'Row count'])
for cell in gs[1]: cell.font = BOLD; cell.fill = HFILL
for (ex, f, why), n in sorted(gap_summary.items(), key=lambda kv: (kv[0][0], kv[0][1], -kv[1])): gs.append([ex, f, why, n])
gs.append([]); gs.append(['TOTAL', '', '', f'=SUM(D2:D{len(gap_summary) + 1})'])
for row in gs.iter_rows(min_row=2):
    for cell in row: cell.font = ARIAL
gs.freeze_panes = 'A2'; gs.auto_filter.ref = f'A1:D{len(gap_summary) + 1}'
for i, w in enumerate([16, 30, 100, 12], 1): gs.column_dimensions[get_column_letter(i)].width = w
gd = wb.create_sheet('Gaps detail'); gd.append(['Exchange', 'Symbol', 'Legal name', 'Security type', 'Field', 'Reason'])
for cell in gd[1]: cell.font = BOLD; cell.fill = HFILL
for g in gap_detail: gd.append(list(g))
for row in gd.iter_rows(min_row=2):
    for cell in row: cell.font = ARIAL
gd.freeze_panes = 'A2'; gd.auto_filter.ref = f'A1:F{len(gap_detail) + 1}'
HITL = [('Zefix official API (ZefixPublicREST)', 'The federal register index is read through the public web API behind zefix.ch (no account; UID, legal seat, legal form, status, SHAB date, cantonal excerpt link). The official ZefixPublicREST needs a free account and answers 401 without one; not needed for the fields carried here.'),
        ('BX Swiss primary listings', 'The instruments list carries every BX-traded line (13,890, mostly sponsored foreign shares and certificates); Swiss share lines not on SIX are carried as BX Swiss rows. A primary-listing flag per company needs the exchange.'),
        ('Annual reports', 'SIX hosts no annual-report index; auditor and report links come from the issuer sites in the Fill pass.'),
        ('Anthropic Admin API key', 'Balance and cost reads for the spend line need an Admin API key; the Messages key answers 401.')]
hs = wb.create_sheet('HITL — needs MK'); hs.append(['Item', 'What is needed'])
for cell in hs[1]: cell.font = BOLD; cell.fill = HFILL
for h in HITL: hs.append(list(h))
for row in hs.iter_rows(min_row=2):
    for cell in row: cell.font = ARIAL; cell.alignment = Alignment(wrap_text=True, vertical='top')
hs.column_dimensions['A'].width = 34; hs.column_dimensions['B'].width = 140
mt = wb.create_sheet('Method'); mt.append(['Item', 'Detail'])
for cell in mt[1]: cell.font = BOLD; cell.fill = HFILL
n_map = sum(1 for r in rows if r.get('isin') and ISIN_LEI.get(r['isin'])); n_exact = sum(1 for r in rows if not (r.get('isin') and ISIN_LEI.get(r['isin'])) and (E_MATCH.get(r['name']) or {}).get('matches')); n_z = sum(1 for z in ZEFIX.values() if z.get('uid'))
METHOD = [('Order', 'ORDER-015 Part A Switzerland Width 0 sweep, node ch-cm-kg, swept ' + TODAY + '. Fly by wire; sourced or blank; no prices or market data carried. Pond-native. Duty line: FinSA.'),
          ('Rosters', "SIX Swiss Exchange: the exchange's public share-explorer feed (PortalSegment EQ: short name, ISIN, valor symbol and number, share type SS/RS/BS/PC, first trading date, number in issue, currency; ICB industry requested but empty in the feed). BX Swiss: the exchange's instruments list, Swiss share lines not on SIX. Lines with a foreign ISIN are kept as exchange lines outside the Swiss corporate scope."),
          ('LEI', f'GLEIF ISIN-to-LEI mapping file, exact ISIN match ({n_map} rows); fallback GLEIF name-exact ({n_exact} rows). Registration status, legal address, headquarters and legal jurisdiction from the LEI record.'),
          ('Zefix (UID)', f'zefix.ch public web API name search per Swiss corporate line; a candidate becomes the entry only when its name equals the roster short name after normalisation and its legal form is a corporation ({n_z} entries). Carried: UID, CH-ID, legal seat, legal form, status, last SHAB date, the cantonal excerpt URL. A UID that disagrees with the LEI record registeredAs is State conflict.'),
          ('Registered office / HQ', 'GLEIF legal address and headquarters city; HQ falls back to the Zefix legal seat.'),
          ('Auditor', 'Not read at Width 0 (source link only per the order); SIX hosts no annual-report index, so the link stays blank with a gap until the Fill pass.'),
          ('Share registrar', 'Tavily page text, agent taken only from a closed list (Computershare Schweiz, areg.ch, ShareCommService, Devigus, SIX SIS, Nimbus, in-house Aktienregister) within 300 characters of "share register" / "Aktienregister"; ties blank and flagged.'),
          ('Newswire of habit', 'Tavily search restricted to PR Newswire, GlobeNewswire, Business Wire, ACCESS Newswire, Newsfile and EQS News; majority of up to three releases. SIX ad hoc publicity is the Width 1 primary.'),
          ('Sources that answered machines', 'six-group.com share-explorer feed (JSON, no key) · bxswiss.com instruments list (HTML) · zefix.ch web API (JSON, no key) · GLEIF API and mapping file · Tavily.'),
          ('Sources that refused machines', 'ZefixPublicREST (401 without an account) · SIX issuer list downloads (404 behind the site) · SER ad hoc pages (404 on the old paths; Width 1 finds the current feed).'),
          ('State', 'sourced = one source read; conflict = the second registry disagrees; filled / confirmed belong to the Fill and Confirm passes.')]
for a, b in METHOD: mt.append([a, b])
for row in mt.iter_rows(min_row=2):
    for cell in row: cell.font = ARIAL; cell.alignment = Alignment(wrap_text=True, vertical='top')
mt.column_dimensions['A'].width = 34; mt.column_dimensions['B'].width = 150
outdir = pond.assembled(NODE); out = os.path.join(outdir, 'ch-issuers.xlsx'); wb.save(out)
os.makedirs(os.path.dirname(MIRROR), exist_ok=True); shutil.copy(out, MIRROR)
print('saved', out, 'mirrored to', MIRROR)
for ex, data in sheets.items():
    print(f'== {ex}: {len(data)} rows (Swiss ISIN {sum(1 for o in data if o["ISIN"].startswith("CH"))})')
    for f in FIELDS:
        n = sum(1 for o in data if o[f]); print(f'   {f:36s} {n:5d} {n / max(1, len(data)) * 100:5.1f}%')
print('gap rows', len(gap_detail))
