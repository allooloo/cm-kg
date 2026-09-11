r"""Merges every worker's events, dedupes on (exchange, ticker, type, date, URL), applies the name rule to search-found wire hits, and writes
CM-KG\DISCLOSURE\uk-disclosure.xlsx (Events, Coverage, Gaps, Method, HITL) and CM-KG\DISCLOSURE\events\uk-events.jsonl."""
import json, os, sys, datetime
from collections import Counter, defaultdict
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter
from common import *
import pond
NODE = 'uk-cm-kg'
OUT_X = sys.argv[1] if len(sys.argv) > 1 else r'C:\ALLOOLOO\CM-KG\DISCLOSURE\uk-disclosure.xlsx'
OUT_J = sys.argv[2] if len(sys.argv) > 2 else r'C:\ALLOOLOO\CM-KG\DISCLOSURE\events\uk-events.jsonl'
issuers = load_issuers(); byk = {key(r): r for r in issuers}; bysym = {r['ticker']: r for r in issuers if r['exchange'] == 'Aquis Stock Exchange'}
ALIASES = {}
_al = pond.latest(NODE, 'width1', 'aliases.json') or ('raw/aliases.json' if os.path.exists('raw/aliases.json') else None)
if _al: ALIASES = json.load(open(_al, encoding='utf-8'))  # written by the Fill pass (ORDER-009): key -> [alias, ...], every alias sourced
def load(fn):
    """pond rule 3: every drop of this rail (POND\\uk-cm-kg\\width1\\<date>\\) plus the working folder; the dedupe key below keeps one copy per event"""
    out = []; seen_paths = set()
    paths = [p for d, p in pond.drops(NODE, 'width1')] + (['raw'] if os.path.isdir('raw') else [])
    for p in paths:
        f = os.path.join(p, os.path.basename(fn)); rp = os.path.realpath(f)
        if not os.path.exists(f) or rp in seen_paths: continue
        seen_paths.add(rp)
        for line in open(f, encoding='utf-8'):
            try: out.append(json.loads(line))
            except Exception: pass
    return out
events = []; gaps = []; worker_notes = defaultdict(list)
def take(rows, worker):
    for d in rows:
        k = d.get('key')
        if d.get('error'): gaps.append((k, worker, 'error: ' + str(d['error'])[:160])); continue
        if d.get('gap') and d['gap'] != 'skipped in this pass': gaps.append((k, worker, d['gap']))
        for e in d.get('events', []): events.append(e)
        if d.get('ambiguous'): worker_notes[k].append(f"{d['ambiguous']} wire hits named another issuer (not taken)")
        if d.get('note'): worker_notes[k].append(d['note'])
take(load('raw/investegate.jsonl'), 'investegate'); take(load('raw/ch_filings.jsonl'), 'ch_filings'); take(load('raw/wire_search.jsonl'), 'wire_search')
for d in load('raw/wire_search.jsonl'):  # search hits without published_date whose URL carries the date
    r = byk.get(d.get('key'))
    if not r: continue
    for u in d.get('undated', []):
        dt = url_date(u['url'])
        if dt and in_window(dt): events.append(event(r, classify(u['title'], default='newswire_release'), dt, u['title'], u['wire'], u['url'], 'Tavily search (wire domains) + URL date', wire=u['wire']))
# Aquis announcements feed (the exchange's own list, read in a browser session; keyed by the exchange symbol)
AQ = pond.latest(NODE, 'width1', 'aquis_announcements.json') or 'raw/aquis_announcements.json'; n_aq = 0
if os.path.exists(AQ):
    aq = json.load(open(AQ, encoding='utf-8'))
    for row in aq.get('rows', []):
        aid, sym, isin, name, title, date = row[:6]
        r = bysym.get(sym)
        if not r or not in_window(date): continue
        events.append(event(r, classify(title), date, title, 'Aquis Stock Exchange announcements', f'https://www.aquis.eu/stock-exchange/announcements/{aid}', 'aquis.eu announcements page data (exchange feed; read in a browser session)', wire='RNS (via Aquis)', rns_category=title)); n_aq += 1
for e in load('raw/grok_live.jsonl'): events.append(e)   # weekly live layer (read by Grok (live)), present on refresh runs only
# pond rule 3: the last versioned event set is merged too (history is never dropped); the dedupe key below keeps one copy
n_prior = 0
_prev = pond.latest_assembled(NODE, 'uk-events.jsonl') or (OUT_J if os.path.exists(OUT_J) else None)
if _prev:
    for line in open(_prev, encoding='utf-8'):
        try:
            e = json.loads(line); e.pop('as_of', None); e.pop('node', None); e.pop('width', None); events.append(e); n_prior += 1
        except Exception: pass
    print('merged prior assembled events', n_prior, 'from', _prev, flush=True)
for e in events: e.pop('key', None)
events = [e for e in events if e.get('url', '').startswith('http') and e.get('date')]
# name rule on search-found wire hits; company pages, the exchange feed and the registry are the issuer's own listings (exempt)
kept = []; ambiguous = Counter()
for e in events:
    if 'company page' in e['read_by'] or 'Companies House' in e['read_by'] or 'aquis.eu' in e['read_by'] or 'Grok' in e['read_by']:
        kept.append(e); continue
    k = e['exchange'] + '|' + e['ticker']
    mode, tok = name_match(e['title'], e['issuer'], tuple(ALIASES.get(k, [])))
    if mode == 'full': kept.append(e)
    elif mode == 'distinctive':
        e['detail'] = (e['detail'] + '; ' if e['detail'] else '') + f"matched on distinctive token '{tok}'"; kept.append(e)
    else:
        ambiguous[k] += 1
        why = (f"token '{tok}' is not unique across the 1,570-row roster or is a common English word" if mode == 'ambiguous' else 'no name token in the title')
        gaps.append((k, 'ambiguous match', f"{e['date']} {e['title'][:90]} | {why} | {e['url']}"))
events = kept
def prio(e): return 0 if 'company page' in e['read_by'] or 'aquis.eu' in e['read_by'] else (1 if 'Companies House' in e['read_by'] else 2)
seen = {}
for e in sorted(events, key=prio):
    kk = (e['exchange'], e['ticker'], e['event_type'], e['date'], e['url'].lower().rstrip('/'))
    if kk not in seen: seen[kk] = e
events = sorted(seen.values(), key=lambda e: (e['exchange'], e['ticker'], e['date'], e['event_type']))
os.makedirs(os.path.dirname(OUT_J), exist_ok=True)
with open(OUT_J, 'w', encoding='utf-8') as f:
    for e in events: f.write(json.dumps({**e, 'as_of': TODAY.isoformat(), 'node': 'uk-cm-kg', 'width': 1}, ensure_ascii=False) + '\n')
CTRL = re.compile(r'[\x00-\x08\x0b\x0c\x0e-\x1f]')
wb = Workbook(); wb.remove(wb.active)
ARIAL = Font(name='Arial', size=10); BOLD = Font(name='Arial', size=10, bold=True); HFILL = PatternFill('solid', fgColor='DDE4EE')
def sheet(name, cols, rows, widths):
    ws = wb.create_sheet(name); ws.append(cols)
    for c in ws[1]: c.font = BOLD; c.fill = HFILL; c.alignment = Alignment(wrap_text=True, vertical='top')
    for r in rows: ws.append([CTRL.sub('', v) if isinstance(v, str) else v for v in r])
    for row in ws.iter_rows(min_row=2):
        for c in row: c.font = ARIAL
    ws.freeze_panes = 'A2'; ws.auto_filter.ref = f'A1:{get_column_letter(len(cols))}{max(len(rows) + 1, 2)}'
    for i, w in enumerate(widths, 1): ws.column_dimensions[get_column_letter(i)].width = w
    return ws
ECOLS = ['Exchange', 'Ticker', 'ISIN', 'LEI', 'Issuer', 'Event type', 'Date', 'Title', 'RNS category', 'CH filing type', 'Wire', 'Source', 'URL', 'Read by', 'Detail', 'State']
erows = [[e['exchange'], e['ticker'], e['isin'], e['lei'], e['issuer'], e['event_type'], e['date'], e['title'], e.get('rns_category', ''), e.get('ch_filing_type', ''), e['wire'], e['source'], e['url'], e['read_by'], e['detail'], e.get('state', 'sourced')] for e in events]
sheet('Events', ECOLS, erows, [16, 10, 15, 22, 34, 22, 11, 60, 30, 20, 16, 30, 60, 40, 30, 10])
per = Counter((e['exchange'], e['ticker']) for e in events)
by_type = defaultdict(Counter)
for e in events: by_type[(e['exchange'], e['ticker'])][e['event_type']] += 1
TYPES = ['regulatory_announcement', 'financial_statement', 'agm_record_date', 'director_dealing', 'early_warning', 'corporate_action', 'halt_resume', 'newswire_release', 'ch_accounts_filed', 'ch_confirmation_statement', 'ch_officer_change', 'ch_name_change', 'ch_charge', 'ch_capital', 'ch_resolution', 'ch_psc', 'ch_other']
CCOLS = ['Exchange', 'Ticker', 'ISIN', 'Issuer', 'Companies House number', 'Newswire of habit', 'Events (live)', 'Events (built)'] + TYPES + ['Zero events (live)', 'Notes']
n_ev = len(events) + 1; crows = []
for r in issuers:
    k = (r['exchange'], r['ticker']); i = len(crows) + 2
    crows.append([r['exchange'], r['ticker'], r['isin'], r['name'], r['ch_number'], r['wire'], f"=COUNTIFS(Events!$A$2:$A${n_ev},A{i},Events!$B$2:$B${n_ev},B{i})", per.get(k, 0)] + [by_type[k].get(t, 0) for t in TYPES] + [f'=IF(G{i}=0,1,0)', '; '.join(worker_notes.get(key(r), []))[:400]])
cv = sheet('Coverage', CCOLS, crows, [16, 10, 15, 34, 14, 18, 12, 12] + [12] * len(TYPES) + [12, 60])
zc = get_column_letter(9 + len(TYPES))
cv.append([]); cv.append(['TOTAL', '', '', '', '', '', f'=SUM(G2:G{len(crows) + 1})', f'=SUM(H2:H{len(crows) + 1})'] + [f'=SUM({get_column_letter(9 + j)}2:{get_column_letter(9 + j)}{len(crows) + 1})' for j in range(len(TYPES))] + [f'=SUM({zc}2:{zc}{len(crows) + 1})', 'issuers with zero events (live)'])
for c in cv[cv.max_row]: c.font = BOLD
GCOLS = ['Exchange', 'Ticker', 'Issuer', 'Worker / field', 'Reason']
grows = []
for k, w, why in gaps:
    r = byk.get(k, {'exchange': k.split('|')[0], 'ticker': k.split('|')[1], 'name': ''}); grows.append([r['exchange'], r['ticker'], r['name'], w, why])
zero = [r for r in issuers if per.get((r['exchange'], r['ticker']), 0) == 0]
for r in zero: grows.append([r['exchange'], r['ticker'], r['name'], 'coverage', 'zero events in the window from every source that answers machines'])
BROKEN = [('all', '', '', 'FCA National Storage Mechanism', 'api.data.fca.org.uk answers HTTP 403; the search UI sits behind a terms-of-use acceptance (HITL). No NSM events.'),
          ('LSE Main Market / AIM', '', '', 'LSE news list and market notices', 'the LSE site loads its news list and notices through component calls the public pages API does not expose (404 on rns-notices / notices); RNS comes from Investegate instead; no LSE market-notice events.'),
          ('Aquis Stock Exchange', '', '', 'Aquis public API', 'public-api.elements.aquis.tech answers HTTP 401 to plain clients; the announcements feed was read through the exchange page data in a browser session (12 months, keyed by symbol).'),
          ('all', '', '', 'Newswire (non-regulatory)', 'GlobeNewswire and Business Wire block direct fetch; wire releases come from Tavily search results, dated from published_date or the release URL.')]
for b in BROKEN: grows.append(list(b))
gs = sheet('Gaps', GCOLS, grows, [18, 10, 34, 24, 110])
gs.append([]); gs.append(['TOTAL gap rows', '', '', '', f'=COUNTA(E2:E{len(grows) + 1})'])
HITL = [('FCA NSM search', 'The NSM UI requires accepting a terms-of-use modal before any search; the search API is not public. A click by MK on data.fca.org.uk (accept terms) would let a browser-session read run; a machine route still needs an FCA data account.'),
        ('Aquis public API token', 'public-api.elements.aquis.tech needs an authentication token the aquis.eu site holds; an Aquis data account (Aquis Elements) would give a machine route instead of the browser-session read.'),
        ('LSE news / notices API', 'The LSE issuer news list and market notices render through authenticated component calls (feedhandler token/saml); an LSEG developer account would be the machine route. Not blocking: Investegate carries RNS.')]
hs = sheet('HITL — needs MK', ['Item', 'What is needed'], [list(h) for h in HITL], [30, 150])
for row in hs.iter_rows(min_row=2):
    for c in row: c.alignment = Alignment(wrap_text=True, vertical='top')
METHOD = [('Order', f'ORDER-008 United Kingdom Width 1: Disclosure. Window {SINCE} to {TODAY}. Scope: {len(issuers)} corporate issuers from uk-issuers.xlsx (security type Corporate; funds, investment companies, depositary receipts and debt out).'),
          ('Rule', 'Blank is blank. No event without a URL. No date read out of a headline: the event date is the announcement, filing or release date. Wire hits naming another issuer are counted in Coverage notes, not taken.'),
          ('Name rule (CEO, 2026-09-10)', 'A search-found wire release is attributed to an issuer only when its title carries every key token of the issuer name (or of a sourced alias), or a single token that is unique across the full 1,570-row roster of uk-issuers.xlsx and not a common English word (common_words.txt). Otherwise Gaps "ambiguous match". Reads from the issuer\'s own Investegate company page, the Companies House filing history and the Aquis announcements feed are the issuer\'s own listings and are exempt.'),
          ('Regulatory announcements (RNS and other RIS)', 'Investegate per-company listing (investegate.co.uk/company/<TIDM>, 50 per page, paged back to the window start): date, RIS source code, RNS headline, announcement URL. rns_category = the RNS headline as published. Event type from the headline (common.py TYPE_RULES): financial_statement (results, trading updates), agm_record_date (AGM/GM notices and results, record dates), director_dealing (PDMR, board changes, adviser changes), early_warning (Holding(s) in Company / TR-1 / Form 8), corporate_action (dividends, placings, issues, consolidations, name changes, admissions, offers, buybacks, total voting rights), halt_resume (suspension, restoration, cancellation), regulatory_announcement (everything else).'),
          ('Companies House filings', 'REST API filing history per issuer with a Companies House number (600 calls per five minutes): filing date, registry category and description, form type; URL = the filing document page. Types ch_accounts_filed · ch_confirmation_statement · ch_officer_change · ch_name_change · ch_charge · ch_capital · ch_resolution · ch_psc · ch_other. Issuers without a number (see uk-issuers Gaps) carry a Gaps line.'),
          ('Exchange notices', f'Aquis: the announcements feed on aquis.eu (exchange page data, {n_aq} rows in the window matched to roster symbols), read in a browser session. LSE market notices and AIM notices: not reachable by machine (see Gaps / HITL).'),
          ('Newswire (search)', 'Tavily news search over prnewswire.com, prnewswire.co.uk, globenewswire.com, businesswire.com, accesswire.com, newsfilecorp.com with published_date; a general query "<short name>" announces on the same domains; date from published_date or the release URL. RNS copies are left to Investegate.'),
          ('Idempotence', 'Dedupe key (exchange, ticker, event type, date, URL). Weekly refresh runs the same workers with WINDOW_DAYS=7 and merges on that key.'),
          ('Read-by labels', 'investegate.co.uk company page (RNS mirror) · Companies House REST API (filing history) · aquis.eu announcements page data (exchange feed; read in a browser session) · Tavily search (wire domains) + published_date · Tavily search (wire domains) + URL date'),
          ('State', 'sourced = one source read. Fill and Confirm (ORDER-009) add filled / confirmed / conflict.')]
mt = sheet('Method', ['Item', 'Detail'], [list(m) for m in METHOD], [30, 150])
for row in mt.iter_rows(min_row=2):
    for c in row: c.alignment = Alignment(wrap_text=True, vertical='top')
os.makedirs(os.path.dirname(OUT_X), exist_ok=True); wb.save(OUT_X)
# versioned output in the pond (rule 2), mirrored at the canonical paths above
_ad = pond.assembled(NODE); import shutil as _sh
_sh.copy(OUT_X, os.path.join(_ad, 'uk-disclosure.xlsx')); _sh.copy(OUT_J, os.path.join(_ad, 'uk-events.jsonl'))
print('saved', OUT_X, 'and', OUT_J, '| versioned copy', _ad)
print('events', len(events)); print('by type', Counter(e['event_type'] for e in events).most_common())
print('by source', Counter(e['read_by'] for e in events).most_common())
print('issuers with events', len(per), '/', len(issuers), '| zero events', len(zero))
print('zero-event first 20:', [f"{r['exchange']}:{r['ticker']} {r['name']}" for r in zero[:20]])
print('gap rows', len(grows), '| by worker', Counter(g[1] for g in gaps).most_common())
print('by exchange', Counter(e['exchange'] for e in events).most_common())
