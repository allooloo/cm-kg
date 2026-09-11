"""Rebuilds CM-KG\ISSUERS\ca-issuers.xlsx with passes 2 and 3 applied:
  new columns 'Also known as' (+ source, read by), per-field '<Field> state' and second-source columns, a summary 'State' column,
  Claude rulings written into ISIN / Auditor where the cell was blank (read by 'adjudicated by Claude'), and a Method tab listing all
  eight engines with the exact duty performed on Canada and the count of fields touched. The Width 0 workbook is kept as
  ca-issuers-width0.xlsx beside it."""
import shutil
from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter
from fc_common import *
from collections import Counter
SRC = ISSUERS_XLSX; BAK = SRC.replace('ca-issuers.xlsx', 'ca-issuers-width0.xlsx')
if not os.path.exists(BAK): shutil.copy(SRC, BAK)
wb = load_workbook(BAK)
ARIAL = Font(name='Arial', size=10); BOLD = Font(name='Arial', size=10, bold=True); HFILL = PatternFill('solid', fgColor='DDE4EE')
conf = {d['key']: d for d in jload('raw/confirm_fields.jsonl')}
aliases = {d['key']: d.get('aliases', []) for d in jload('raw/aliases.jsonl')}
for k_, a_ in perplexity_aliases().items(): aliases.setdefault(k_, []).append(a_)
for k_, a_ in website_aliases().items(): aliases.setdefault(k_, []).append(a_)
claude_isin = {d['key']: d for d in jload('raw/claude_isin.jsonl')}; claude_aud = {d['key']: d for d in jload('raw/claude_auditor.jsonl')}
rulings = {d['key']: d for d in jload('raw/conflict_rulings.jsonl')}
FIELDS = ['ISIN', 'LEI', 'Transfer agent', 'Auditor', 'Newswire of habit', 'Incorporation jurisdiction', 'Also known as']
touched = Counter(); states = Counter(); gaps_new = []
for ex in ['TSX', 'TSXV', 'CSE', 'Cboe Canada']:
    ws = wb[ex]; hdr = [c.value for c in ws[1]]; H = {h: i + 1 for i, h in enumerate(hdr)}
    base = len(hdr)
    new = ['Also known as', 'Alias source', 'Alias read by'] + [f'{f} state' for f in FIELDS] + [f'{f} second source' for f in FIELDS] + [f'{f} second read by' for f in FIELDS] + ['State']
    for i, h in enumerate(new, base + 1):
        c = ws.cell(row=1, column=i, value=h); c.font = BOLD; c.fill = HFILL; c.alignment = Alignment(wrap_text=True, vertical='top')
    col = {h: base + 1 + i for i, h in enumerate(new)}
    for rr in range(2, ws.max_row + 1):
        k = ex + '|' + str(ws.cell(row=rr, column=H['Ticker']).value); c = conf.get(k, {}).get('fields', {})
        # Claude fills
        ci = claude_isin.get(k, {})
        if ci.get('ruling') and not ws.cell(row=rr, column=H['ISIN']).value:
            ws.cell(row=rr, column=H['ISIN'], value=ci['ruling']); ws.cell(row=rr, column=H['ISIN source'], value=ci.get('source') or ws.cell(row=rr, column=H['ISIN source']).value); ws.cell(row=rr, column=H['ISIN read by'], value='adjudicated by Claude'); touched['Claude:ISIN'] += 1
            c['ISIN'] = {'state': 'filled', 'second': [ci.get('source', '')], 'read_by2': 'adjudicated by Claude'}
        elif k in claude_isin and not ci.get('ruling'):
            gaps_new.append((ex, k.split('|')[1], 'ISIN', 'adjudicated by Claude: left blank — ' + (ci.get('reason') or '')[:160]))
            g = ws.cell(row=rr, column=H['Gaps']); g.value = (g.value or '') + f"; ISIN: adjudicated by Claude, left blank: {(ci.get('reason') or '')[:120]}"
        ca = claude_aud.get(k, {})
        if ca.get('ruling') and not ws.cell(row=rr, column=H['Auditor']).value:
            ws.cell(row=rr, column=H['Auditor'], value=ca['ruling']); ws.cell(row=rr, column=H['Auditor source'], value=ca.get('source') or ''); ws.cell(row=rr, column=H['Auditor read by'], value='adjudicated by Claude'); touched['Claude:Auditor'] += 1
            c['Auditor'] = {**c.get('Auditor', {}), 'state': c.get('Auditor', {}).get('state') if c.get('Auditor', {}).get('state') in ('confirmed', 'conflict') else 'filled', 'read_by2': c.get('Auditor', {}).get('read_by2') or 'adjudicated by Claude'}
        elif k in claude_aud and not ca.get('ruling'):
            gaps_new.append((ex, k.split('|')[1], 'Auditor', 'adjudicated by Claude: left blank — ' + (ca.get('reason') or '')[:160]))
        # Pass 3 conflict rulings (adjudicated by Claude): a ruling stands as 'filled'; an unsettled conflict blanks the cell with the reason in Gaps
        for f, (colv, colsrc) in {'Transfer agent': ('Transfer agent', 'Transfer agent source'), 'Auditor': ('Auditor', 'Auditor source'), 'ISIN': ('ISIN', 'ISIN source'), 'Incorporation jurisdiction': ('Incorporation jurisdiction', 'Jurisdiction source'), 'Newswire of habit': ('Newswire of habit', 'Newswire releases seen')}.items():
            ru = rulings.get(k + '#' + f)
            if not ru or f not in c: continue
            rbcol = {'Transfer agent': 'Transfer agent read by', 'Auditor': 'Auditor read by', 'ISIN': 'ISIN read by', 'Incorporation jurisdiction': 'Jurisdiction read by', 'Newswire of habit': 'Newswire read by'}[f]
            if ru.get('ruling'):
                ws.cell(row=rr, column=H[colv], value=ru['ruling']); ws.cell(row=rr, column=H[rbcol], value=(ws.cell(row=rr, column=H[rbcol]).value or '') + ' · adjudicated by Claude'); touched['Claude:conflicts ruled'] += 1
                c[f] = {**c[f], 'state': 'filled', 'read_by2': f"adjudicated by Claude (reading {ru.get('basis', '')}): {ru.get('reason', '')[:120]}"}
            else:
                ws.cell(row=rr, column=H[colv], value=None); touched['Claude:conflicts blanked'] += 1
                gaps_new.append((ex, k.split('|')[1], f, f"conflict unresolved, cell blanked — A: {ru.get('first', '')[:60]} vs B: {'; '.join(str(x) for x in ru.get('second', []))[:80]} — adjudicated by Claude: {ru.get('reason', '')[:120]}"))
                c[f] = {**c[f], 'state': 'conflict', 'read_by2': f"adjudicated by Claude: unsettled — {ru.get('reason', '')[:120]}"}
        # aliases
        al = aliases.get(k, [])
        if al:
            ws.cell(row=rr, column=col['Also known as'], value=' | '.join(a['alias'] for a in al)); ws.cell(row=rr, column=col['Alias source'], value='\n'.join(a['source'] for a in al)); ws.cell(row=rr, column=col['Alias read by'], value='\n'.join(a['read_by'] for a in al))
            for a in al: touched[a['read_by'].split(' (')[0].split(' —')[0]] += 1
        # states
        summary = Counter()
        for f in FIELDS:
            v = ws.cell(row=rr, column=H[f]).value if f in H else (ws.cell(row=rr, column=col['Also known as']).value if f == 'Also known as' else None)
            st = c.get(f, {}).get('state') if v else ''
            if v and not st: st = 'filled' if 'adjudicated' in str(ws.cell(row=rr, column=H[f]).value if f in H else '') else 'sourced'
            if st:
                ws.cell(row=rr, column=col[f'{f} state'], value=st); ws.cell(row=rr, column=col[f'{f} second source'], value='\n'.join(x for x in c.get(f, {}).get('second', []) if x)); ws.cell(row=rr, column=col[f'{f} second read by'], value=c.get(f, {}).get('read_by2', ''))
                summary[st] += 1; states[(f, st)] += 1
        ws.cell(row=rr, column=col['State'], value=' · '.join(f'{s} {n}' for s, n in sorted(summary.items())) if summary else '')
    for rr in range(2, ws.max_row + 1):
        for i in range(base + 1, base + 1 + len(new)): ws.cell(row=rr, column=i).font = ARIAL
    for h, i in col.items(): ws.column_dimensions[get_column_letter(i)].width = 30 if 'source' in h or h == 'Also known as' else 18
    ws.auto_filter.ref = f'A1:{get_column_letter(ws.max_column)}{ws.max_row}'
# Gaps additions
gs = wb['Gaps'];
for row in gaps_new: gs.append(list(row) + [1])
for r_ in gs.iter_rows(min_row=gs.max_row - len(gaps_new), max_row=gs.max_row):
    for c_ in r_: c_.font = ARIAL
# Method tab: eight engines
def method_rows():
    a_g = sum(1 for k, l in aliases.items() for a in l if a['read_by'].startswith('GLEIF'))
    a_w = sum(1 for k, l in aliases.items() for a in l if 'company page' in a['read_by'] and 'Perplexity' not in a['read_by'])
    a_p = sum(1 for k, l in aliases.items() for a in l if 'Perplexity' in a['read_by'])
    a_x = sum(1 for k, l in aliases.items() for a in l if a['read_by'].startswith('exchange') or a['read_by'].startswith('thecse'))
    st = Counter(); [st.__setitem__(s, st[s] + n) for (f, s), n in states.items()]
    cf = sum(1 for d in conf.values() for f, v in d['fields'].items() if v.get('state') in ('confirmed', 'conflict') and 'Tavily' in v.get('read_by2', ''))
    return [
        ('Claude (claude-sonnet-5)', f"Adjudicated Width 0 conflicts from the cited sources: ISIN {len(claude_isin)} cases ({sum(1 for d in claude_isin.values() if d.get('ruling'))} ruled, {sum(1 for d in claude_isin.values() if not d.get('ruling'))} left blank with reason); auditor {len(claude_aud)} cases ({sum(1 for d in claude_aud.values() if d.get('ruling'))} ruled, {sum(1 for d in claude_aud.values() if not d.get('ruling'))} left blank). Pass 3 conflicts (first source vs second source): {len(rulings)} cases, {sum(1 for d in rulings.values() if d.get('ruling'))} ruled, {sum(1 for d in rulings.values() if not d.get('ruling'))} unsettled and blanked. Fields written to this workbook: {touched['Claude:ISIN'] + touched['Claude:Auditor'] + touched['Claude:conflicts ruled']} written, {touched['Claude:conflicts blanked']} blanked. Label 'adjudicated by Claude'."),
        ('ChatGPT (gpt-4.1-mini, Batch API)', 'Touched nothing on the issuer record. Its duty was the financial_statement extraction on the Disclosure workbook (see that Method tab).'),
        ('Google Gemini (gemini-flash-latest)', 'Touched nothing on the issuer record. Its duty was the AGM / record-date read on the Disclosure workbook.'),
        ('Perplexity (sonar-pro, grounded search)', f"Found the current wire company page for issuers whose page was missing or stale in Width 1: {sum(1 for d in jload('raw/perplexity_pages.jsonl') if d.get('verified'))} verified of {len(jload('raw/perplexity_pages.jsonl'))} asked; {a_p} aliases written from the page display names. Label 'read by Perplexity (agent)'. The /v1/agent endpoint exists but was not used; grounded chat completions were."),
        ('Grok (grok-4.6, Responses API + web/X search)', 'Touched nothing on the issuer record. Wired into the daily refresh as the 48-hour live layer for halts, resumes and newswire (grok_live.py); validated on one TSXV query on 2026-09-10, no events written for the backfill.'),
        ('Mistral (mistral-small-latest)', 'Touched nothing on the issuer record. Its duty was the French-release read on the Disclosure workbook.'),
        ('Tavily (search)', f"Second-source reads for Pass 3: transfer agent, auditor and ISIN confirmations — {cf} fields resolved to confirmed or conflict from a second page on a different domain; alias rematch searches for zero-event and ambiguous issuers (see Disclosure)."),
        ('Cloudflare', 'Touched nothing on the issuer record (network and site only).'),
        ('Sources (not engines)', f"Aliases: GLEIF other entity names {a_g}; wire company pages {a_w}; exchange profile names {a_x}. States across the seven fields: " + ', '.join(f'{s} {n}' for s, n in sorted(st.items())) + '.'),
        ('Rule', 'Aliases sourced only (GLEIF, the issuer\'s own wire company page, the exchange profile); never a guess. States: sourced = one source read; filled = written by adjudication or a lab read; confirmed = two independent sources agree, both cited; conflict = the second source disagrees, sent to Claude or left blank with reason.'),
    ]
if 'Method' in wb.sheetnames:
    mt = wb['Method']; mt.append([]); mt.append(['Passes 2 and 3 (ORDER-004, ' + TODAY.isoformat() + ')', 'Engine duty and count of fields touched on Canada'])
    for c in mt[mt.max_row]: c.font = BOLD
    for a, b in method_rows():
        mt.append([a, b])
        for c in mt[mt.max_row]: c.font = ARIAL; c.alignment = Alignment(wrap_text=True, vertical='top')
wb.save(SRC)
import pond; _ad = pond.assembled('ca-cm-kg'); shutil.copy(SRC, os.path.join(_ad, 'ca-issuers.xlsx')); print('versioned copy', _ad)  # pond rule 2
print('saved', SRC); print('touched', dict(touched)); print('states', sorted(states.items()))
