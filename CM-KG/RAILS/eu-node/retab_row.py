r"""retab_row.py — DAY ORDER 5 (CEO tab rule, 2026-10-24): the ROW workbooks re-issued with the tab names of record ('Country — Stock Exchange' / 'Country — CODE',
RAILS\eu-node\tabs.json). Pond rule: the assembled drop of record is never edited — a new assembled drop is written (pond.assembled numbers it above the day's
highest), the same rows and events, sheets renamed only; the Exchange column keeps the official name. Mirrored to CM-KG\ISSUERS (a .pre copy beside, as eu_node
does). The nightly jobs re-assemble on the new rule from their next image. Usage: python -X utf8 retab_row.py <cc> ..."""
import os, sys, json, shutil, datetime
import openpyxl
HERE = os.path.dirname(os.path.abspath(__file__)); sys.path.insert(0, HERE); import pond
ROOT = r'C:\ALLOOLOO'
TABS = json.load(open(os.path.join(HERE, 'tabs.json'), encoding='utf-8')); VEN = json.load(open(os.path.join(HERE, 'venues.json'), encoding='utf-8'))
for cc in sys.argv[1:]:
    node = f'{cc}-cm-kg'; xp = pond.latest_assembled(node, f'{cc}-issuers.xlsx')
    if not xp: print(cc, 'no assembled workbook'); continue
    src = os.path.dirname(xp); wb = openpyxl.load_workbook(xp); ren = []
    for mic, name in VEN[cc].items():
        new = TABS[cc][mic][:31]
        for old in (name[:31], name):
            if old in wb.sheetnames and old != new: wb[old].title = new; ren.append((old, new)); break
    if not ren: print(cc, 'already on the rule', wb.sheetnames); continue
    out = pond.assembled(node); wb.save(os.path.join(out, f'{cc}-issuers.xlsx'))
    for f in os.listdir(src):
        if f != f'{cc}-issuers.xlsx': shutil.copyfile(os.path.join(src, f), os.path.join(out, f))
    bj = os.path.join(out, 'build.json'); b = json.load(open(bj)) if os.path.exists(bj) else {}
    b['retab'] = {'from': os.path.basename(src), 'renamed': ren, 'rule': TABS['_rule'], 'at': datetime.datetime.now(datetime.timezone.utc).isoformat()}; json.dump(b, open(bj, 'w'), indent=1, ensure_ascii=False)
    dst = os.path.join(ROOT, 'CM-KG', 'ISSUERS', f'{cc}-issuers.xlsx')
    if os.path.exists(dst): shutil.copyfile(dst, dst + '.pre-' + datetime.datetime.now().strftime('%Y%m%dT%H%M%S'))
    shutil.copyfile(os.path.join(out, f'{cc}-issuers.xlsx'), dst)
    print(cc, os.path.basename(src), '->', os.path.basename(out), ren)
