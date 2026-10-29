r"""hk-width1 assembler (ORDER-035, 2026-10-22). Reads every hkexnews-<board>-<day>.json from every pond drop of this rail (newest drop wins per board-day),
keeps the documents of the 2,820 roster lines (Main Board, GEM, REITs — warrants, CBBCs, ETPs and debt are not on the roster and are not carried),
English titles only (a title carrying a Chinese character is skipped and counted, never read), the trailing twelve months, and writes
  POND\hk-cm-kg\assembled\<date>\hk-events.jsonl + hk-disclosure.xlsx   (versioned, immutable)
  CM-KG\DISCLOSURE\events\hk-events.jsonl + CM-KG\DISCLOSURE\hk-disclosure.xlsx   (the mirror the door loader reads; the previous mirror is kept beside as .pre-<stamp>)
Event rule: type = the HKEX headline category as the exchange's own code table writes it, in English (hkex_category), plus the house event_type mapped
from it; title = exactly as filed (line breaks joined by a space), never translated; identifiers (stock code, ISIN, LEI) language-free.
Dedupe on (stock code, hkex_category, date, url). Source and read date on every event."""
import os, re, sys, json, glob, html, datetime, shutil
from collections import Counter, defaultdict
import openpyxl
from openpyxl import Workbook
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__))); import pond
NODE = 'hk-cm-kg'; BASE = 'https://www1.hkexnews.hk'
ISSUERS = pond.latest_assembled(NODE, 'hk-issuers.xlsx') or r'C:\ALLOOLOO\CM-KG\ISSUERS\hk-issuers.xlsx'
AS_OF = os.environ.get('AS_OF_HK') or os.environ.get('SWEEP_MARKET_DATE') or datetime.date.today().isoformat()
CJK = re.compile(r'[\u3000-\u303f\u3400-\u4dbf\u4e00-\u9fff\uf900-\ufaff\uff00-\uffef]')
TABS = {'Main Board': 'HKEX', 'GEM': 'HKEX-GEM', 'REITs': 'HKEX-REIT'}
def roster():
    wb = openpyxl.load_workbook(ISSUERS, read_only=True); out = {}
    for tab, ex in TABS.items():
        ws = wb[tab]; it = ws.iter_rows(values_only=True); hdr = next(it)
        for r in it:
            d = dict(zip(hdr, r))
            if d.get('Stock code'): out[str(d['Stock code']).zfill(5)] = {'tab': tab, 'exchange': ex, 'name': d.get('Name of Securities (HKEX, English)'), 'isin': d.get('ISIN') or '', 'lei': d.get('LEI') or ''}
    return out
def house_type(cat):
    c = cat.lower()
    if 'annual report' in c and 'cancelled' not in c: return 'annual_report'   # the combined heading [Annual Report / ESG ...] is the annual report too
    if 'interim/half-year report' in c or '[interim report' in c: return 'interim_report'
    if 'esg' in c and 'financial statements' in c: return 'esg_report'
    if 'results' in c and ('final results' in c or 'interim results' in c or 'quarterly results' in c or 'results of' not in c and 'financial statements' in c): return 'results'
    if 'trading halt' in c or 'suspension' in c or 'resumption' in c: return 'halt_suspension'
    if 'change in directors' in c or 'important executive functions' in c: return 'director_change'
    if 'dividend' in c or 'distribution' in c and 'design' not in c: return 'dividend'
    if 'change of company name' in c: return 'name_change'
    if 'results of agm' in c or 'results of egm' in c or 'notice of agm' in c or 'notice of egm' in c or 'proxy forms' in c or 'meeting' in c: return 'agm_egm'
    if c.startswith('monthly returns'): return 'monthly_return'
    if c.startswith('next day disclosure returns'): return 'next_day_disclosure'
    if c.startswith('circulars'): return 'circular'
    if c.startswith('listing documents'): return 'listing_document'
    if c.startswith('constitutional documents'): return 'constitutional_document'
    if c.startswith('documents on display'): return 'documents_on_display'
    if 'takeovers code' in c: return 'takeovers'
    if 'substantial shareholder' in c or 'disclosure of interests' in c: return 'substantial_shareholder'
    return 'announcement'
def clean(s): return re.sub(r'\s+', ' ', html.unescape(html.unescape(re.sub(r'<[^>]+>', ' ', str(s or ''))))).strip()
if __name__ == '__main__':
    R = roster(); files = {}
    for date, p in pond.drops(NODE, 'width1'):
        for f in glob.glob(os.path.join(p, 'hkexnews-*-*.json')): files[os.path.basename(f)] = f   # newest drop wins per board-day
    today = datetime.date.fromisoformat(AS_OF[:10]); start = today - datetime.timedelta(days=366)
    ev = {}; st = Counter(); cats = Counter(); per = Counter(); incomplete = []
    for name, f in sorted(files.items()):
        d = json.load(open(f, encoding='utf-8'))
        if not d.get('complete'): incomplete.append((d['board'], d['date'], d['rows_returned'], d['total_count']))
        for row in d['rows']:
            st['rows'] += 1
            dt = datetime.datetime.strptime(row['DATE_TIME'], '%d/%m/%Y %H:%M'); day = dt.date()
            if not (start <= day <= today): st['outside_window'] += 1; continue
            title = clean(row.get('TITLE')); cat = clean(row.get('LONG_TEXT'))
            if CJK.search(title) or CJK.search(cat): st['skipped_chinese'] += 1; continue
            url = BASE + row['FILE_LINK'] if row.get('FILE_LINK', '').startswith('/') else row.get('FILE_LINK')
            codes = [c.zfill(5) for c in re.findall(r'\d{1,5}', clean(row.get('STOCK_CODE')).replace(' ', '')) if c]
            hit = False
            for code in codes:
                r = R.get(code)
                if not r: continue
                hit = True; k = (code, cat, day.isoformat(), url)
                if k in ev: st['dupes'] += 1; continue
                ev[k] = {'exchange': r['tab'], 'exchange_code': r['exchange'], 'ticker': code, 'isin': r['isin'], 'lei': r['lei'], 'issuer': r['name'], 'event_type': house_type(cat), 'date': day.isoformat(),
                         'time_hkt': dt.strftime('%H:%M'), 'title': title, 'hkex_category': cat, 'wire': 'HKEXnews', 'source': 'HKEXnews (listed company information, title search)', 'url': url,
                         'read_by': 'HKEXnews title search (titleSearchServlet.do, lang=E)', 'reference': row.get('NEWS_ID', ''), 'file_type': row.get('FILE_TYPE') or 'HTML', 'file_info': row.get('FILE_INFO', ''),
                         'state': 'sourced', 'read_at': d['read_at'][:10], 'as_of': AS_OF[:10], 'node': NODE, 'width': 1}
                cats[cat] += 1; per[code] += 1
            if not hit: st['not_on_roster'] += 1
    events = sorted(ev.values(), key=lambda e: (e['ticker'], e['date'], e['time_hkt']), reverse=True)
    stamp = datetime.datetime.now().strftime('%Y%m%dT%H%M%S'); adir = pond.assembled(NODE)
    def write_jsonl(p):
        with open(p, 'w', encoding='utf-8') as f:
            for e in events: f.write(json.dumps(e, ensure_ascii=False) + '\n')
    write_jsonl(os.path.join(adir, 'hk-events.jsonl'))
    wb = Workbook(); ws = wb.active; ws.title = 'Events'; cols = ['exchange', 'exchange_code', 'ticker', 'issuer', 'isin', 'event_type', 'hkex_category', 'date', 'time_hkt', 'title', 'url', 'source', 'read_by', 'state', 'reference']
    ws.append(cols); [ws.append([e.get(c) for c in cols]) for e in events]
    cv = wb.create_sheet('Coverage'); cv.append(['Measure', 'Value'])
    for k, v in [('events', len(events)), ('issuers with at least one event', len(per)), ('roster lines', len(R)), ('issuers with zero events', len(R) - len(per)), ('board-days read', len(files)), ('board-days incomplete', len(incomplete)),
                 ('rows read', st['rows']), ('skipped: Chinese-language title', st['skipped_chinese']), ('skipped: not on the roster (warrants, CBBCs, ETPs, debt, delisted)', st['not_on_roster']), ('outside the twelve-month window', st['outside_window']), ('duplicates', st['dupes'])]: cv.append([k, v])
    ct = wb.create_sheet('Categories'); ct.append(['HKEX headline category (English, as the exchange writes it)', 'events', 'house event_type'])
    for c, n in cats.most_common(): ct.append([c, n, house_type(c)])
    me = wb.create_sheet('Method'); me.append(['Item', 'Detail'])
    me.append(['Order', 'ORDER-035 HONG KONG NODE (2026-10-22): Width 1 from HKEXnews public pages; English-language filings only (standing rule); financialfilings.com never a source (terms).'])
    me.append(['Source', 'HKEXnews title search JSON (titleSearchServlet.do, lang=E), one call per board per day; the document link, category, title and time as HKEXnews lists them.'])
    me.append(['Rule', 'Type in English from the HKEX headline category table; title exactly as filed; no translation written to the record; a Chinese-language title is skipped and counted.'])
    if incomplete: me.append(['Incomplete board-days', json.dumps(incomplete)[:30000]])
    wb.save(os.path.join(adir, 'hk-disclosure.xlsx'))
    for src, dst in ((os.path.join(adir, 'hk-events.jsonl'), r'C:\ALLOOLOO\CM-KG\DISCLOSURE\events\hk-events.jsonl'), (os.path.join(adir, 'hk-disclosure.xlsx'), r'C:\ALLOOLOO\CM-KG\DISCLOSURE\hk-disclosure.xlsx')):
        os.makedirs(os.path.dirname(dst), exist_ok=True)   # DAY ORDER 4 (2026-10-24): the job image carries no DISCLOSURE folder — the first two hk nightlies failed here
        if os.path.exists(dst): shutil.copyfile(dst, f'{dst}.pre-{stamp}')   # copyfile: the image's path shim does not cover the xattr calls of copy2
        shutil.copyfile(src, dst)
    print(json.dumps({'assembled': adir, 'events': len(events), 'issuers_with_events': len(per), 'roster': len(R), 'stats': st, 'incomplete_board_days': len(incomplete), 'top_categories': cats.most_common(8)}, ensure_ascii=False, default=str))
