"""roster_v13.py — SITE-04: roster v1.3, the workbook of record, from v1.2: ruling (5) data-lifecycle.ai → role 301 → data-vintage.ai (estate allooloo, lens Shelf);
the ROSTER NOTE — the 13 EU-node issuer hostnames (at be cz dk es fi gr ie it no pl pt se -issuer.ai) do not resolve in DNS and are not zones on the account:
each gets a roster row with role 301 → eu-issuer.ai, written on the wire once the zone exists (the CEO's hand: the domain). v1.2 stays beside, untouched.
Usage: python -X utf8 -I roster_v13.py"""
import openpyxl, datetime, shutil
from openpyxl.styles import Font
SRC = r'C:\ALLOOLOO\TRAFFIC\ALLOOLOO_TRAFFIC_zone-roster_v1.2_2026-10-28.xlsx'; OUT = r'C:\ALLOOLOO\TRAFFIC\ALLOOLOO_TRAFFIC_zone-roster_v1.3_2026-10-28.xlsx'
wb = openpyxl.load_workbook(SRC); ws = wb['ZONES']; hdr = [c.value for c in ws[1]]; col = {h: i + 1 for i, h in enumerate(hdr)}; F = Font(name='Arial')
EU13 = ['at', 'be', 'cz', 'dk', 'es', 'fi', 'gr', 'ie', 'it', 'no', 'pl', 'pt', 'se']
done = []
for r in range(2, ws.max_row + 1):
    if ws.cell(r, col['zone']).value == 'data-lifecycle.ai':
        ws.cell(r, col['role']).value = '301'; ws.cell(r, col['target']).value = 'https://data-vintage.ai/'
        ws.cell(r, col['role_source']).value = 'Ruling (5), CEO 2026-10-28, applied at SITE-04: role 301 → data-vintage.ai, estate allooloo, lens Shelf'
        ws.cell(r, col['tag_of_record']).value = 'estate:allooloo lens:Shelf family:shelf role:301 target:data-vintage.ai site-04:2026-10-28 ruling:5'
        ws.cell(r, col['ruling']).value = '(5) 301 → data-vintage.ai'; ws.cell(r, col['writes']).value = (ws.cell(r, col['writes']).value or '') + '; SITE-04: AAAA 100:: apex+www, redirect rule 301 → data-vintage.ai, tag role:301 (ruling5_apply.json)'
        done.append('data-lifecycle.ai')
for cc in EU13:
    row = {h: None for h in hdr}
    row.update({'zone': f'{cc}-issuer.ai', 'estate': 'allooloo', 'family': 'issuer twin (EU-node member)', 'lens': 'KYP', 'include': 'N', 'role': '301', 'target': 'https://eu-issuer.ai/',
                'role_source': 'ROSTER NOTE, CEO 2026-10-28 (SITE-04): hostname does not resolve in DNS and is not a zone on the account; role 301 → eu-issuer.ai of record, written on the wire once the zone exists (the domain: the CEO\'s hand). The SM-ECO-10060 member pages already link eu-issuer.ai (ruling c).',
                'wire_today_status': 'no DNS', 'wire_today_final': None, '301_check': 'not on the account', 'tag_of_record': 'estate:allooloo lens:KYP family:issuer twin (EU-node member) role:301 target:eu-issuer.ai site-04:2026-10-28 (not yet written: no zone)',
                'ai_bots_protection': None, 'dnssec': None, 'requests_30d': 0, 'zone_id': None, 'created_on': None, 'ruling': 'ROSTER NOTE 2026-10-28', 'writes': 'none (no zone)'})
    ws.append([row[h] for h in hdr])
    for c in ws[ws.max_row]: c.font = F
    done.append(f'{cc}-issuer.ai')
ws.auto_filter.ref = ws.dimensions
cfg = wb['CONFIG']
for r in range(1, cfg.max_row + 1):
    if cfg.cell(r, 1).value == 'roster':
        cfg.cell(r, 2).value = 'v1.3 — SITE-04 (2026-10-28): v1.2 + ruling (5) data-lifecycle.ai → 301 → data-vintage.ai (Shelf) + the ROSTER NOTE: 13 EU-node issuer hostnames (not on the account, no DNS) carry role 301 → eu-issuer.ai pending the zone. Read nightly by the standing TRAFFIC run (cmkg-traffic-nightly) for estate · family · lens · include.'
cfg.append(['prior', 'ALLOOLOO_TRAFFIC_zone-roster_v1.2_2026-10-28.xlsx (SITE-03), kept beside']); cfg.append(['nightly run', 'cmkg-traffic-nightly (Allooloo tenant, allooloo-cmkg-eastus) 02:30 UTC: C:\\ALLOOLOO\\TRAFFIC\\rails\\traffic_nightly.py → TRAFFIC-<date>.xlsx + traffic-<date>.json + latest.json in allooloocmkgapex/ops/traffic/; the digest line reads latest.json; the RADAR tile stays the typed LENS-01 figure (CEO, 2026-10-28)'])
for r in cfg.iter_rows(min_row=cfg.max_row - 1):
    for c in r: c.font = F
wb.save(OUT); print('wrote', OUT, 'rows changed/added', done)
