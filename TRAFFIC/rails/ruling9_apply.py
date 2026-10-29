"""ruling9_apply.py — Ruling (9), CEO 2026-10-28: standard-10060.* (all TLDs) → estate allooloo, lens Proof; supersedes the gsc default of ruling (4).
Roster v1.3 rows (standard-10060.org serve; .com .ca .info 301 → standard-10060.org, the edge 308 of record) and the TXT _estate tag of record on each zone
(SITE-03 pattern; the token cannot write Cloudflare zone tags). standard-10060-dpp.* are another name and stay as they are. Usage: python -X utf8 -I ruling9_apply.py [--apply]"""
import json, sys, urllib.request, urllib.error, openpyxl
from openpyxl.styles import Font
T = open(r'C:\ALLOOLOO\AGENT KEYS\cloudflare.txt').read().strip(); H = {'Authorization': 'Bearer ' + T, 'Content-Type': 'application/json'}; B = 'https://api.cloudflare.com/client/v4'; APPLY = '--apply' in sys.argv
def call(method, path, body=None):
    try: return json.load(urllib.request.urlopen(urllib.request.Request(B + path, data=json.dumps(body).encode() if body is not None else None, headers=H, method=method), timeout=60))
    except urllib.error.HTTPError as e: return json.load(e)
X = r'C:\ALLOOLOO\TRAFFIC\ALLOOLOO_TRAFFIC_zone-roster_v1.3_2026-10-28.xlsx'; wb = openpyxl.load_workbook(X); ws = wb['ZONES']; hdr = [c.value for c in ws[1]]; col = {h: i + 1 for i, h in enumerate(hdr)}
log = []
for r in range(2, ws.max_row + 1):
    z = ws.cell(r, col['zone']).value
    if not (isinstance(z, str) and z.startswith('standard-10060.')): continue
    serve = z == 'standard-10060.org'; zid = ws.cell(r, col['zone_id']).value
    tag = f"estate:allooloo lens:Proof family:standard-10060 role:{'serve' if serve else '301'}{'' if serve else ' target:standard-10060.org'} std-01:2026-10-28 ruling:9"
    ws.cell(r, col['estate']).value = 'allooloo'; ws.cell(r, col['family']).value = 'standard-10060'; ws.cell(r, col['lens']).value = 'Proof'; ws.cell(r, col['include']).value = 'Y'
    ws.cell(r, col['role']).value = 'serve' if serve else '301'; ws.cell(r, col['target']).value = None if serve else 'https://standard-10060.org/'
    ws.cell(r, col['role_source']).value = 'Ruling (9), CEO 2026-10-28: standard-10060.* → estate allooloo, lens Proof (supersedes the gsc default of ruling 4); ' + ('the canonical page (Pages project cull-canonical-core-org-ai)' if serve else 'the edge 308 → standard-10060.org of record')
    ws.cell(r, col['tag_of_record']).value = tag; ws.cell(r, col['ruling']).value = '(9) standard-10060.* → allooloo / Proof'
    res = None
    if zid:
        dns = call('GET', f'/zones/{zid}/dns_records?type=TXT&name=_estate.{z}').get('result') or []
        if APPLY: res = (call('PATCH', f'/zones/{zid}/dns_records/{dns[0]["id"]}', {'content': tag}) if dns else call('POST', f'/zones/{zid}/dns_records', {'type': 'TXT', 'name': '_estate.' + z, 'content': tag, 'ttl': 1, 'comment': 'STD-01 / ruling (9) (2026-10-28): estate and lens tag of record'})).get('success')
    ws.cell(r, col['writes']).value = ((ws.cell(r, col['writes']).value or '') + f"; ruling 9: tag TXT _estate = {res if APPLY else 'dry'}").strip('; ')
    for c in ws[r]: c.font = Font(name='Arial')
    log.append((z, 'serve' if serve else '301', res if APPLY else 'dry'))
cfg = wb['CONFIG']; cfg.append(['ruling (9)', 'CEO 2026-10-28: standard-10060.* (all TLDs) → estate allooloo, lens Proof; supersedes the gsc default of ruling (4). standard-10060-dpp.* are another name, unchanged (gsc).'])
for c in cfg[cfg.max_row]: c.font = Font(name='Arial')
if APPLY: wb.save(X)
print('applied' if APPLY else 'dry run', log)
