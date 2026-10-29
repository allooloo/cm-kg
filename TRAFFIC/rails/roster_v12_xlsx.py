"""SITE-03 — roster v1.2, the workbook of record, rewritten from the plan and the apply/verify logs (the LENS-01 v1.2 stays beside as .pre).
Sheets: CONFIG · ZONES (every zone: estate · family · lens · role · target · role_source · wire today · 301 check · tags · ai_bots · dnssec) · FORWARDS (the dormant stock moved)
· FAMILIES · LENS (zones + 30-day requests per group, incl. Forwards; sm-eco-10060.* on its own line) · TAGS (counts per lens group) · APPLY (every write and its result).
Usage: python -X utf8 -I roster_v12_xlsx.py <plan.json> <apply_log.json> <verify.json> <traffic_raw.json> <out.xlsx>"""
import json, sys, os, shutil, datetime
from collections import Counter, defaultdict
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
P = json.load(open(sys.argv[1], encoding='utf-8')); A = {z['zone']: z for z in json.load(open(sys.argv[2], encoding='utf-8'))['zones']}; V = {z['zone']: z for z in json.load(open(sys.argv[3], encoding='utf-8'))['zones']}
T = json.load(open(sys.argv[4], encoding='utf-8')); req = {zid: sum(g['sum']['requests'] for g in gs) for zid, gs in T['daily'].items()}
OUT = sys.argv[5]
if os.path.exists(OUT): shutil.copy2(OUT, OUT.replace('.xlsx', '.pre-site03-20261028.xlsx'))
HF = Font(name='Arial', bold=True, color='FFFFFF'); HFILL = PatternFill('solid', fgColor='1F3864'); F = Font(name='Arial'); BF = Font(name='Arial', bold=True)
wb = openpyxl.Workbook(); wb.remove(wb.active)
def sheet(name, header, rows, widths=None, freeze='A2'):
    ws = wb.create_sheet(name); ws.append(header)
    for c in ws[1]: c.font = HF; c.fill = HFILL; c.alignment = Alignment(vertical='center', wrap_text=True)
    for r in rows: ws.append(r)
    for row in ws.iter_rows(min_row=2):
        for c in row: c.font = F
    for i, h in enumerate(header, 1): ws.column_dimensions[get_column_letter(i)].width = (widths or {}).get(h, max(12, min(60, len(str(h)) + 4)))
    if freeze: ws.freeze_panes = freeze
    ws.auto_filter.ref = ws.dimensions; return ws
acts = lambda n: '; '.join(f"{a['action']} = {'ok' if a['ok'] else ('dry' if a['ok'] is None else 'FAIL ' + a['detail'])}" for a in (A.get(n) or {}).get('actions', []))
cfg = [['key', 'value'], ['roster', 'v1.2 — SITE-03 (Site Build Plan v2.2 §5/§8, 2026-10-28): the four LENS-01 rulings applied; the role of record per Allooloo zone; every §5 301 written as a Cloudflare redirect rule; DORMANT STOCK re-pointed; AI-bot block OFF and DNSSEC ON; the estate/lens tag of record on every Allooloo zone'],
       ['account', 'GSC and Allooloo AI Agent Stack dd2832b36f171b815f84c8487aada36b (not split; one lens by tag)'], ['input', 'LENS-01 roster v1.2 (2026-10-28) + wire read + Cloudflare state read (RAILS wire_read.py, api_state.py, rulings_v12.py, merge_site03.py, apply_site03.py, verify_site03.py)'],
       ['rulings', '(1) 413 distinct is the roster of record; (2) the 24 roster-gap zones: country issuer pages → KYP, do-eco/hk-eco → ECO, agentic-x402 → Proof, agent-swap → hold; (3) agentic-trades → Shelf, agentmarketnews → allooloo / Issuerwire; (4) sm-eco-10060.* → allooloo / ECO, acm-* → allooloo / Proof, sm-cpg-10060/sm-aio-cpg/sm-amp-cpg/standard-10060 → gsc'],
       ['tags', 'Cloudflare zone tags need a token scope this token lacks (Zone Tags, error 10000) — the CEO creates it. Until then the tag of record is the TXT record _estate.<zone> ("estate:… lens:… family:… role:…"), written on every Allooloo zone, and the estate/lens columns here. The lens reads either.'],
       ['lens groups', 'Company · KYP · Issuerwire · Shelf · ECO · Proof · Forwards (allooloo-forward: the dormant stock, its own line, never mixed into KYP or Issuerwire); hold = agent-swap (CEO to name); Excluded = agentic-zen/agenticzen'],
       ['window', 'traffic column: LENS-01 read, 2026-09-28 .. 2026-10-27 UTC'], ['written', datetime.datetime.now(datetime.UTC).strftime('%Y-%m-%dT%H:%M:%SZ') + ' by CC, SITE-03']]
ws = wb.create_sheet('CONFIG'); [ws.append(r) for r in cfg]
for c in ws['A']: c.font = BF
for r in ws.iter_rows(min_row=1): r[1].font = F; r[1].alignment = Alignment(wrap_text=True, vertical='top')
ws.column_dimensions['A'].width = 16; ws.column_dimensions['B'].width = 160
order = {'allooloo': 0, 'allooloo-forward': 1, 'office': 2, 'exclude': 3, 'gsc': 4}
zh = ['zone', 'estate', 'family', 'lens', 'include', 'role', 'target', 'role_source', 'wire_today_status', 'wire_today_final', '301_check', 'tag_of_record', 'ai_bots_protection', 'dnssec', 'requests_30d', 'zone_id', 'created_on', 'ruling', 'writes']
zr = []
for z in sorted(P, key=lambda z: (order[z['estate']], z.get('lens', ''), z['family'], z['zone'])):
    v = V.get(z['zone']) or {}; chk = ('' if z['estate'] not in ('allooloo', 'allooloo-forward') else ('ok' if v.get('ok') else (f"FAIL {v.get('status')} → {v.get('location', '')[:40]} / target {v.get('target_status')}" if v else 'not checked')))
    zr.append([z['zone'], z['estate'], z['family'], z.get('lens', ''), z.get('include', ''), z.get('role', ''), z.get('target', ''), z.get('role_source', ''), str(z.get('wire_status', '')), z.get('wire_final', '')[:80], chk, z.get('tag_txt', ''), str(z.get('ai_bots_protection', '')), str(z.get('dnssec', '')), req.get(z['zone_id']), z['zone_id'], z.get('created_on', ''), z.get('ruling', ''), acts(z['zone'])])
sheet('ZONES', zh, zr, {'zone': 32, 'family': 24, 'lens': 20, 'target': 40, 'role_source': 48, 'wire_today_final': 44, '301_check': 30, 'tag_of_record': 60, 'zone_id': 34, 'ruling': 40, 'writes': 80})
D = [z for z in P if z['estate'] == 'allooloo-forward']
sheet('FORWARDS', ['zone', 'inventory family', 'why', 'target', 'wire before', '301 check', 'writes'], [[z['zone'], z['family'], z['dormant_kind'], z['target'], z.get('wire_final', '')[:60], ('ok' if (V.get(z['zone']) or {}).get('ok') else 'FAIL'), acts(z['zone'])] for z in sorted(D, key=lambda z: z['zone'])], {'zone': 32, 'why': 24, 'wire before': 44, 'writes': 80})
fam = Counter((z['estate'], z['family']) for z in P); fr = [[e, f, c, sum(req.get(z['zone_id'], 0) for z in P if z['estate'] == e and z['family'] == f)] for (e, f), c in sorted(fam.items())]
sheet('FAMILIES', ['estate', 'family', 'zones', 'requests_30d'], fr, {'family': 36})
lens = defaultdict(lambda: [0, 0])
for z in P:
    if z['estate'] in ('allooloo', 'allooloo-forward'): lens[z['lens']][0] += 1; lens[z['lens']][1] += req.get(z['zone_id'], 0)
sm = [z for z in P if z['family'] == 'sm-eco-10060']
six = ['Company', 'KYP', 'Issuerwire', 'Shelf', 'ECO', 'Proof']
lr = [[k, lens[k][0], lens[k][1], 'in the estate number' if k in six else ('its own line — in the Allooloo lens, never mixed into KYP or Issuerwire' if k == 'Forwards' else 'beside it')] for k in six + ['Forwards'] + [k for k in lens if k not in six and k != 'Forwards']]
lr.append(['ESTATE (six groups)', sum(lens[k][0] for k in six), sum(lens[k][1] for k in six), 'the one number']); lr.append(['ESTATE + Forwards', sum(lens[k][0] for k in six) + lens['Forwards'][0], sum(lens[k][1] for k in six) + lens['Forwards'][1], ''])
lr.append(['of which sm-eco-10060.* (apex + members, every TLD)', len(sm), sum(req.get(z['zone_id'], 0) for z in sm), 'reported as its own line at SITE-03 close (ruling 4); counted inside ECO'])
sheet('LENS', ['lens group', 'zones', 'requests_30d (LENS-01 window)', 'counted'], lr, {'lens group': 48, 'counted': 70}, freeze=None)
tags = Counter(z.get('lens') for z in P if z.get('tag_txt'))
sheet('TAGS', ['lens group', 'zones tagged (TXT _estate of record)', 'estate tag'], [[k, v, 'allooloo-forward' if k == 'Forwards' else 'allooloo'] for k, v in sorted(tags.items())], {'lens group': 40, 'zones tagged (TXT _estate of record)': 36}, freeze=None)
ar = [[z['zone'], a['action'], 'ok' if a['ok'] else ('dry' if a['ok'] is None else 'FAIL'), a['detail'][:200]] for z in A.values() for a in z['actions']]
sheet('APPLY', ['zone', 'write', 'result', 'detail'], ar, {'zone': 32, 'write': 50, 'detail': 90})
wb.save(OUT); print('wrote', OUT, len(zr), 'zones;', len(D), 'forwards;', len(ar), 'writes')
