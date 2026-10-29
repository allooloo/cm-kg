"""SITE-03B — crawl-state-<date>.xlsx: the table read before anything changed (and, after the apply, the AFTER sheet). Usage: python -X utf8 -I crawl_xlsx.py <crawl_state.json> <out.xlsx> [<apply_log.json> <after.json> <proof.json>]"""
import json, sys, datetime
from collections import Counter
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
C = json.load(open(sys.argv[1], encoding='utf-8')); OUT = sys.argv[2]
A = json.load(open(sys.argv[3], encoding='utf-8')) if len(sys.argv) > 3 else None; AF = json.load(open(sys.argv[4], encoding='utf-8')) if len(sys.argv) > 4 else None; PR = json.load(open(sys.argv[5], encoding='utf-8')) if len(sys.argv) > 5 else None
HF = Font(name='Arial', bold=True, color='FFFFFF'); HFILL = PatternFill('solid', fgColor='1F3864'); F = Font(name='Arial'); BF = Font(name='Arial', bold=True)
wb = openpyxl.Workbook(); wb.remove(wb.active)
def sheet(name, header, rows, widths=None):
    ws = wb.create_sheet(name); ws.append(header)
    for c in ws[1]: c.font = HF; c.fill = HFILL; c.alignment = Alignment(vertical='center', wrap_text=True)
    for r in rows: ws.append(r)
    for row in ws.iter_rows(min_row=2):
        for c in row: c.font = F
    for i, h in enumerate(header, 1): ws.column_dimensions[get_column_letter(i)].width = (widths or {}).get(h, max(12, min(60, len(str(h)) + 4)))
    ws.freeze_panes = 'A2'; ws.auto_filter.ref = ws.dimensions; return ws
def rows_of(zs):
    out = []
    for o in sorted(zs, key=lambda o: ({'allooloo': 0, 'allooloo-forward': 1, 'office': 2, 'exclude': 3, 'gsc': 4}.get(o['estate'], 5), o['zone'])):
        out.append([o['zone'], o['estate'], o['plan'], str(o.get('ai_bots_protection')), str(o.get('crawler_protection')), str(o.get('content_bots_protection')), str(o.get('fight_mode')), str(o.get('sbfm_verified_bots')), str(o.get('sbfm_definitely_automated')), str(o.get('sbfm_likely_automated')), str(o.get('is_robots_txt_managed')),
                    o.get('waf_custom_rules', 0), '; '.join(f"{r['action']}{'' if r['enabled'] else ' (disabled)'}: {r['expression'][:80]}" for r in o.get('waf_custom_bot_rules', [])), '; '.join(f"{r['override_action']}: {r['description']}" for r in o.get('waf_managed_challenge_overrides', [])),
                    str(o.get('robots_status')), ', '.join(o.get('robots_ai_disallowed', [])), '; '.join(f'{u}: {p}' for u, p in o.get('robots_disallow', [])[:12]), 'yes' if o.get('robots_cf_managed') else 'no', 'yes' if o.get('robots_sitemap') else 'no', o.get('robots_served_by', ''), str(o.get('llms_status')), 'yes' if o.get('llms_present') else 'no', o.get('bm_error', '')])
    return out
H = ['zone', 'estate', 'plan', 'ai_bots_protection (AI Crawl Control)', 'crawler_protection', 'content_bots_protection', 'fight_mode (Bot Fight Mode)', 'sbfm_verified_bots', 'sbfm_definitely_automated', 'sbfm_likely_automated', 'is_robots_txt_managed', 'waf custom rules', 'waf custom rules on bots (challenge/block)', 'waf managed challenge overrides', 'robots.txt status', 'AI user agents disallowed', 'Disallow lines (first 12)', 'Cloudflare managed robots block', 'sitemap line', 'robots served by', 'llms.txt status', 'llms.txt present', 'read error']
W = {'zone': 32, 'waf custom rules on bots (challenge/block)': 60, 'Disallow lines (first 12)': 60, 'robots served by': 30, 'AI user agents disallowed': 30}
sheet('BEFORE', H, rows_of(C['zones']), W)
def summary(zs, title):
    rows = []
    for est in ('allooloo', 'allooloo-forward', 'gsc', 'office', 'exclude'):
        s = [o for o in zs if o['estate'] == est]
        if not s: continue
        rows.append([est, len(s), sum(1 for o in s if o.get('ai_bots_protection') == 'block'), sum(1 for o in s if o.get('ai_bots_protection') == 'disabled'), sum(1 for o in s if o.get('fight_mode') is True), sum(1 for o in s if str(o.get('sbfm_verified_bots')) not in ('None', 'allow')), sum(1 for o in s if o.get('is_robots_txt_managed') is True), sum(1 for o in s if o.get('waf_custom_bot_rules')), sum(1 for o in s if o.get('waf_managed_challenge_overrides')), sum(1 for o in s if o.get('robots_status') == 200), sum(1 for o in s if o.get('robots_ai_disallowed')), sum(1 for o in s if o.get('robots_cf_managed')), sum(1 for o in s if o.get('robots_status') == 200 and not o.get('robots_sitemap')), sum(1 for o in s if o.get('llms_present')), sum(1 for o in s if not o.get('bm_readable', True))])
    rows.append(['TOTAL', len(zs), sum(r[2] for r in rows), sum(r[3] for r in rows), sum(r[4] for r in rows), sum(r[5] for r in rows), sum(r[6] for r in rows), sum(r[7] for r in rows), sum(r[8] for r in rows), sum(r[9] for r in rows), sum(r[10] for r in rows), sum(r[11] for r in rows), sum(r[12] for r in rows), sum(r[13] for r in rows), sum(r[14] for r in rows)])
    sheet(title, ['estate', 'zones', 'AI bots: block', 'AI bots: disabled (allowed)', 'Bot Fight Mode on', 'SBFM verified bots not allow', 'managed robots.txt on', 'zones with WAF bot rules', 'zones with managed challenge overrides', 'robots.txt 200', 'robots.txt disallowing AI agents', 'Cloudflare managed robots block', 'robots.txt 200 without sitemap line', 'llms.txt present', 'bot_management unreadable'], rows, {'estate': 20})
summary(C['zones'], 'SUMMARY-BEFORE')
if A:
    ar = [[z['zone'], z.get('estate', ''), a['action'], 'ok' if a['ok'] else ('dry' if a['ok'] is None else 'FAIL'), a['detail'][:200]] for z in A['zones'] for a in z['actions']]
    sheet('CHANGES', ['zone', 'estate', 'write', 'result', 'detail'], ar, {'zone': 32, 'write': 56, 'detail': 80})
if AF: sheet('AFTER', H, rows_of(AF['zones']), W); summary(AF['zones'], 'SUMMARY-AFTER')
if PR: sheet('PROOF', ['zone', 'estate', 'user agent', 'status', 'challenge?', 'bytes', 'title / note'], [[p['zone'], p['estate'], p['ua'], str(p['status']), 'yes' if p.get('challenge') else 'no', p.get('bytes', ''), p.get('note', '')[:120]] for p in PR['rows']], {'zone': 32, 'title / note': 60})
cfg = wb.create_sheet('CONFIG', 0); cfg.append(['key', 'value'])
for k, v in [['order', 'SITE-03B — open to training crawlers, every zone (CEO, 2026-10-28)'], ['read_at', C['read_at']], ['scope', f"{len(C['zones'])} zones on the account dd2832b3… (Allooloo estate, forwards, GSC estate, office, excluded)"], ['sources', 'Cloudflare API: zones/{id}/bot_management · rulesets/phases/http_request_firewall_custom/entrypoint · rulesets/phases/http_request_firewall_managed/entrypoint; the wire: https://<zone>/robots.txt and /llms.txt'], ['SITE-03 value', 'SITE-03 wrote ai_bots_protection = "disabled" (AI bots allowed) on 326 zones that read "block"; nothing to reverse'], ['written', datetime.datetime.now(datetime.UTC).strftime('%Y-%m-%dT%H:%M:%SZ') + ' by CC']]: cfg.append([k, v])
for c in cfg['A']: c.font = BF
for r in cfg.iter_rows(min_row=1): r[1].font = F; r[1].alignment = Alignment(wrap_text=True, vertical='top')
cfg.column_dimensions['A'].width = 16; cfg.column_dimensions['B'].width = 140
wb.save(OUT); print('wrote', OUT, len(C['zones']), 'zones', 'sheets', wb.sheetnames)
