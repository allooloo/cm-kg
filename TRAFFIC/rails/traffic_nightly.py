r"""traffic_nightly.py — SITE-04 (Site Build Plan v2.2 §8, 2026-10-28): the LENS-01 read as the standing TRAFFIC run. Nightly, read-only on Cloudflare:
every zone on the account, GraphQL httpRequests1dGroups for the prior 30 days (daily, UTC), classified by the roster of record (estate · family · lens · include)
with the TXT _estate tag as the fallback for a zone the roster does not know; the ESTATE number = the sum of the six Allooloo lens groups
(Company · KYP · Issuerwire · Shelf · ECO · Proof). Writes TRAFFIC-<date>.xlsx + traffic-<date>.json (+ latest.json, the pointer the nightly digest line reads)
to the private ops store allooloocmkgapex/ops/traffic/ (dated files beside, never over) and, on the laptop, to C:\ALLOOLOO\TRAFFIC\nightly\. Never typed values.
The RADAR ESTATE tile is NOT touched (CEO, 2026-10-28: it stays the typed LENS-01 figure). Token: CLOUDFLARE_ANALYTICS_TOKEN (vault reference in the job) or the
AGENT KEYS file on the laptop, read in-process, never printed. Usage: python -X utf8 -I traffic_nightly.py [--date YYYY-MM-DD] [--no-upload]"""
import os, sys, re, json, time, datetime, threading, urllib.request, urllib.error, concurrent.futures as cf
from collections import Counter, defaultdict
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
ACCT = os.environ.get('CLOUDFLARE_ACCOUNT_ID', 'dd2832b36f171b815f84c8487aada36b')
T = os.environ.get('CLOUDFLARE_ANALYTICS_TOKEN') or open(r'C:\ALLOOLOO\AGENT KEYS\cloudflare-analytics.txt', encoding='utf-8').read().strip()
H = {'Authorization': 'Bearer ' + T, 'Content-Type': 'application/json'}; B = 'https://api.cloudflare.com/client/v4'
ROSTER = os.environ.get('ROSTER_XLSX') or r'C:\ALLOOLOO\TRAFFIC\ALLOOLOO_TRAFFIC_zone-roster_v1.3_2026-10-28.xlsx'
today = datetime.date.fromisoformat(sys.argv[sys.argv.index('--date') + 1]) if '--date' in sys.argv else datetime.datetime.now(datetime.timezone.utc).date()
DATE = today.isoformat(); d_geq = (today - datetime.timedelta(days=30)).isoformat(); d_lt = DATE   # 30 full UTC days ending yesterday
OUT = os.environ.get('OUT_DIR') or (os.path.join(r'C:\ALLOOLOO\TRAFFIC\nightly', DATE) if os.name == 'nt' else f'/tmp/traffic/{DATE}'); os.makedirs(OUT, exist_ok=True)
SIX = ['Company', 'KYP', 'Issuerwire', 'Shelf', 'ECO', 'Proof']
def log(*a): print(datetime.datetime.now(datetime.timezone.utc).strftime('%H:%M:%S'), *a, flush=True)
def api(path):
    for attempt in range(5):
        try: return json.load(urllib.request.urlopen(urllib.request.Request(B + path, headers=H), timeout=120))
        except Exception as e:
            err = repr(e)[:200]; time.sleep(5 * (attempt + 1))
    raise RuntimeError(err)
# 1. every zone on the account
zones = []; page = 1
while True:
    r = api(f'/zones?per_page=50&page={page}&account.id={ACCT}'); zones += r['result']
    if page >= r['result_info']['total_pages']: break
    page += 1
log('zones', len(zones))
# 2. the roster of record
wb = openpyxl.load_workbook(ROSTER, read_only=True); ws = wb['ZONES']; hdr = [c.value for c in next(ws.iter_rows(max_row=1))]
rost = {}
for row in ws.iter_rows(min_row=2, values_only=True):
    d = dict(zip(hdr, row))
    if d.get('zone'): rost[d['zone']] = d
wb.close(); log('roster', os.path.basename(ROSTER), len(rost), 'rows')
def tag_fallback(z):
    """a zone the roster does not know: the TXT _estate tag of record, if the token can read DNS; else unclassified"""
    try:
        r = api(f"/zones/{z['id']}/dns_records?type=TXT&name=_estate.{z['name']}")
        for d in r.get('result') or []:
            kv = dict(p.split(':', 1) for p in d['content'].split() if ':' in p)
            if 'estate' in kv: return kv.get('estate'), kv.get('family', ''), kv.get('lens', ''), 'tag TXT _estate'
    except Exception: pass
    return 'unclassified', '', 'unclassified (new zone — roster + tag say nothing)', 'none'
rows = {}
for z in zones:
    d = rost.get(z['name'])
    if d: rows[z['id']] = {'zone': z['name'], 'estate': d.get('estate'), 'family': d.get('family'), 'lens': d.get('lens'), 'include': d.get('include'), 'source': 'roster'}
    else:
        est, fam, lens, src = tag_fallback(z); rows[z['id']] = {'zone': z['name'], 'estate': est, 'family': fam, 'lens': lens, 'include': 'Y' if est == 'allooloo' else 'N', 'source': src}
new_zones = [r['zone'] for r in rows.values() if r['source'] != 'roster']; log('zones not in the roster', len(new_zones), new_zones[:10])
# 3. the traffic read (GraphQL, 10 zones a query, throttled under the 300 / 5 min quota)
Q = '''query($tags:[string!],$geq:Date!,$lt:Date!){viewer{zones(filter:{zoneTag_in:$tags}){zoneTag
 httpRequests1dGroups(limit:31,filter:{date_geq:$geq,date_lt:$lt},orderBy:[date_ASC]){dimensions{date} sum{requests cachedRequests bytes threats pageViews} uniq{uniques}}}}}'''
lock = threading.Lock(); last = [0.0]; GAP = 1.1; daily = {}; failed = {}
def gql(tags):
    body = json.dumps({'query': Q, 'variables': {'tags': tags, 'geq': d_geq, 'lt': d_lt}}).encode(); err = ''
    for attempt in range(5):
        with lock:
            w = last[0] + GAP - time.time()
            if w > 0: time.sleep(w)
            last[0] = time.time()
        try:
            r = json.load(urllib.request.urlopen(urllib.request.Request(B + '/graphql', data=body, headers=H), timeout=120))
            if r.get('errors'): raise RuntimeError(json.dumps(r['errors'])[:300])
            return r['data']['viewer']['zones']
        except Exception as e: err = repr(e)[:300]; log('gql retry', attempt + 1, err[:120]); time.sleep(10 * (attempt + 1))
    raise RuntimeError(err)
done = [0]
def work(batch):
    try:
        got = {z['zoneTag']: z['httpRequests1dGroups'] for z in gql(batch)}
        with lock:
            for t in batch: daily[t] = got.get(t, [])
    except Exception as e:
        with lock:
            for t in batch: failed[t] = repr(e)[:200]
        log('batch failed:', repr(e)[:160])
    with lock:
        done[0] += 1
        if done[0] % 20 == 0: log('read', done[0], 'batches;', len(daily), 'zones;', len(failed), 'failed')
ids = [z['id'] for z in zones]
with cf.ThreadPoolExecutor(int(os.environ.get('GQL_THREADS', '4'))) as ex: list(ex.map(work, [ids[i:i + 10] for i in range(0, len(ids), 10)]))
log('read', len(daily), 'zones;', len(failed), 'failed')
# 4. the numbers
req = {t: sum(g['sum']['requests'] for g in gs) for t, gs in daily.items()}; byt = {t: sum(g['sum']['bytes'] for g in gs) for t, gs in daily.items()}
days = sorted({g['dimensions']['date'] for gs in daily.values() for g in gs})
groups = Counter(); gz = Counter(); est_tot = Counter(); est_z = Counter()
for t, r in rows.items():
    n = req.get(t, 0); est_tot[r['estate'] or '?'] += n; est_z[r['estate'] or '?'] += 1
    if r['estate'] == 'allooloo' and r['include'] == 'Y' and r['lens'] in SIX: groups[r['lens']] += n; gz[r['lens']] += 1
estate = sum(groups[l] for l in SIX); account = sum(req.values())
unassigned = sorted(r['zone'] for r in rows.values() if r['estate'] == 'allooloo' and r['include'] == 'Y' and r['lens'] not in SIX)
j = {'estate_requests_30d': estate, 'window': [days[0], days[-1]] if days else [d_geq, d_lt], 'read_at': datetime.datetime.now(datetime.timezone.utc).strftime('%Y-%m-%dT%H:%M:%SZ'), 'date': DATE,
     'groups': {l: groups[l] for l in SIX}, 'zones_in_groups': sum(gz[l] for l in SIX), 'account_requests_30d': account, 'share_of_account': round(estate / account, 4) if account else None,
     'zones_read': len(daily), 'zones_failed': len(failed), 'zones_not_in_roster': new_zones, 'allooloo_unassigned_zones': unassigned, 'by_estate': {k: {'zones': est_z[k], 'requests_30d': v} for k, v in est_tot.items()},
     'roster': os.path.basename(ROSTER), 'source': 'SITE-04 standing TRAFFIC run: Cloudflare GraphQL Analytics httpRequests1dGroups, prior 30 days daily UTC, every zone on the account; the six Allooloo lens groups summed (roster of record + TXT _estate tag fallback). The RADAR tile is not written by this run.',
     'written_by': os.environ.get('CONTAINER_APP_JOB_EXECUTION_NAME') or 'laptop', 'image': os.environ.get('IMAGE_TAG')}
json.dump(j, open(os.path.join(OUT, f'traffic-{DATE}.json'), 'w', encoding='utf-8'), indent=1)
# 5. the workbook (values from the read; totals as formulas, recalculated by Excel on open)
HF = Font(name='Arial', bold=True, color='FFFFFF'); HFILL = PatternFill('solid', fgColor='1F3864'); F = Font(name='Arial')
xb = openpyxl.Workbook(); xb.remove(xb.active)
def sheet(name, header, rr, widths=None, freeze='A2'):
    s = xb.create_sheet(name); s.append(header)
    for c in s[1]: c.font = HF; c.fill = HFILL; c.alignment = Alignment(vertical='center', wrap_text=True)
    for r in rr: s.append(r)
    for row in s.iter_rows(min_row=2):
        for c in row: c.font = F
    for i, h in enumerate(header, 1): s.column_dimensions[get_column_letter(i)].width = (widths or {}).get(h, max(12, min(60, len(str(h)) + 4)))
    if freeze: s.freeze_panes = freeze
    s.auto_filter.ref = s.dimensions; return s
sheet('CONFIG', ['key', 'value'], [['run', f'SITE-04 standing TRAFFIC run · {DATE} · written by {j["written_by"]}'], ['window', f'{j["window"][0]} .. {j["window"][1]} UTC (30 days)'], ['roster', j['roster']], ['token', 'read-only Analytics token, by reference, never printed'],
       ['estate number', 'ESTATE!B2 = SUMIFS over the six Allooloo lens groups (Company, KYP, Issuerwire, Shelf, ECO, Proof), include = Y'], ['RADAR tile', 'not written by this run (CEO, 2026-10-28: the tile keeps the typed LENS-01 figure)'], ['digest', 'the nightly digest line reads ops/traffic/latest.json']], {'key': 18, 'value': 120}, freeze=None)
zr = sorted(([r['zone'], r['estate'], r['family'], r['lens'], r['include'], req.get(t, 0), byt.get(t, 0), sum(1 for g in daily.get(t, []) if g['sum']['requests']), r['source'], t] for t, r in rows.items()), key=lambda x: (-x[5], x[0]))
sheet('ZONES', ['zone', 'estate', 'family', 'lens', 'include', 'requests_30d', 'bytes_30d', 'days_with_traffic', 'classified_by', 'zone_id'], zr, {'zone': 36, 'family': 28, 'lens': 22, 'zone_id': 34})
N = len(zr) + 1; RQ = f'ZONES!$F$2:$F${N}'; LN = f'ZONES!$D$2:$D${N}'; ES = f'ZONES!$B$2:$B${N}'; IN = f'ZONES!$E$2:$E${N}'
sheet('BY-LENS', ['lens group', 'zones', 'requests_30d', 'counted'], [[l, gz[l], f'=SUMIFS({RQ},{LN},"{l}",{ES},"allooloo",{IN},"Y")', 'yes'] for l in SIX] + [[f'unassigned ({len(unassigned)} zones)', len(unassigned), sum(req.get(t, 0) for t, r in rows.items() if r['zone'] in unassigned), 'no']], freeze=None)
sheet('ESTATE', ['measure', 'value', 'note'], [['ESTATE (Allooloo, six lens groups, 30-day requests)', '=' + '+'.join(f'SUMIFS({RQ},{LN},"{l}",{ES},"allooloo",{IN},"Y")' for l in SIX), f'{j["window"][0]} .. {j["window"][1]} UTC · value as read: {estate:,}'],
      ['Account, every zone', f'=SUM({RQ})', f'as read: {account:,}'], ['Allooloo share of the account', '=IF(B3=0,0,B2/B3)', f'as read: {j["share_of_account"]}'], ['zones read / failed', f'{len(daily)} / {len(failed)}', '; '.join(f'{rows[t]["zone"]}: {e[:60]}' for t, e in list(failed.items())[:10])]], {'measure': 58, 'value': 24, 'note': 90}, freeze=None)
xp = os.path.join(OUT, f'TRAFFIC-{DATE}.xlsx'); xb.save(xp); log('wrote', xp, 'ESTATE', f'{estate:,}', 'share', j['share_of_account'])
# 6. the private ops store (dated files beside, never over; latest.json is the pointer and is rewritten)
if '--no-upload' not in sys.argv and os.environ.get('OPS_ACCOUNT'):
    from azure.storage.blob import BlobServiceClient, ContentSettings
    if os.environ.get('OPS_KEY'): cred = os.environ['OPS_KEY']   # laptop: the account key fetched in-process by the wrapper (the signed-in user has no data-plane role on the store)
    else:
        from azure.identity import DefaultAzureCredential; cred = DefaultAzureCredential()   # the job: its system identity (Storage Blob Data Contributor on ops)
    c = BlobServiceClient(f"https://{os.environ['OPS_ACCOUNT']}.blob.core.windows.net", credential=cred).get_container_client(os.environ.get('OPS_CONTAINER', 'ops'))
    def beside(name, data, ct):
        n, i = name, 1
        while c.get_blob_client(n).exists(): i += 1; n = re.sub(r'(\.[a-z]+)$', f'-{i}\\1', name)
        c.upload_blob(n, data, content_settings=ContentSettings(content_type=ct)); return n
    n1 = beside(f'traffic/{DATE}/TRAFFIC-{DATE}.xlsx', open(xp, 'rb').read(), 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    n2 = beside(f'traffic/{DATE}/traffic-{DATE}.json', json.dumps(j, ensure_ascii=False, indent=1).encode('utf-8'), 'application/json')
    c.upload_blob('traffic/latest.json', json.dumps(j, ensure_ascii=False, indent=1).encode('utf-8'), overwrite=True, content_settings=ContentSettings(content_type='application/json'))
    log('ops ->', n1, n2, 'traffic/latest.json')
log('done', DATE, 'ESTATE', f'{estate:,}')
