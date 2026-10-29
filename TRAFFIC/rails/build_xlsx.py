"""LENS-01 — write ZONES-2026-10-28.xlsx (the account, measured) and the roster v1.2. Values from the API reads; totals as formulas; recalc after."""
import json, datetime, sys
from collections import defaultdict, OrderedDict
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
SP = r'C:\Users\matth\AppData\Local\Temp\claude\C--bpc-agent\4fc1172b-aaae-4859-9f0f-29f4a875d910\scratchpad'
OUT = r'C:\ALLOOLOO\TRAFFIC'; DATE = '2026-10-28'
Z = json.load(open(f'{SP}/zones_classified.json', encoding='utf-8'))
A = json.load(open(f'{SP}/zones_all.json', encoding='utf-8'))
T = json.load(open(f'{SP}/traffic_raw.json', encoding='utf-8'))
R11 = json.load(open(f'{SP}/roster_v11.json', encoding='utf-8'))
days = [(datetime.date(2026, 9, 28) + datetime.timedelta(days=i)).isoformat() for i in range(30)]
assert days[0] == T['window'][0] and days[-1] < T['window'][1]
HF = Font(name='Arial', bold=True, color='FFFFFF'); HFILL = PatternFill('solid', fgColor='1F3864'); F = Font(name='Arial'); BF = Font(name='Arial', bold=True)
def sheet(wb, name, header, rows, widths=None, freeze='A2'):
    ws = wb.create_sheet(name); ws.append(header)
    for c in ws[1]: c.font = HF; c.fill = HFILL; c.alignment = Alignment(vertical='center', wrap_text=True)
    for r in rows: ws.append(r)
    for row in ws.iter_rows(min_row=2):
        for c in row: c.font = F
    for i, h in enumerate(header, 1): ws.column_dimensions[get_column_letter(i)].width = (widths or {}).get(h, max(12, min(60, len(str(h)) + 4)))
    ws.freeze_panes = freeze; ws.auto_filter.ref = ws.dimensions
    return ws
def buckets(sm):
    b = {'2xx': 0, '3xx': 0, '4xx': 0, '5xx': 0, 'other': 0}
    for e in sm.get('responseStatusMap') or []:
        s = int(e['edgeResponseStatus']); k = f'{s // 100}xx' if 200 <= s < 600 else 'other'; b[k] += e['requests']
    return b
# ---- per zone / per day
by_id = {z['zone_id']: z for z in Z}
daily_rows = []; zone_tot = {}
for zid, groups in T['daily'].items():
    z = by_id[zid]; tot = defaultdict(int); tot['days_with_traffic'] = 0
    for g in groups:
        sm = g['sum']; b = buckets(sm); d = g['dimensions']['date']
        daily_rows.append([z['zone'], d, z['estate'], z['family'], z['lens'], sm['requests'], sm['cachedRequests'], sm['bytes'], sm['cachedBytes'], sm['threats'], sm['pageViews'], sm['encryptedRequests'], g['uniq']['uniques'], b['2xx'], b['3xx'], b['4xx'], b['5xx'], b['other']])
        for k in ('requests', 'cachedRequests', 'bytes', 'cachedBytes', 'threats', 'pageViews', 'encryptedRequests'): tot[k] += sm[k]
        for k, v in b.items(): tot[k] += v
        tot['uniques_daily_sum'] += g['uniq']['uniques']; tot['days_with_traffic'] += 1 if sm['requests'] else 0
    zone_tot[zid] = tot
daily_rows.sort(key=lambda r: (r[0], r[1]))
def top_countries(zid, n=5):
    cm = sorted(T['country'].get(zid) or [], key=lambda e: -e['requests'])[:n]; return ' · '.join(f"{e['clientCountryName']} {e['requests']:,}" for e in cm)
wb = openpyxl.Workbook(); wb.remove(wb.active)
# ---- CONFIG
cfg = [['key', 'value'],
       ['order', 'LENS-01 · The account, measured (Site Build Plan v2.1 §8, 2026-10-28) — read-only; nothing on the wire changed'],
       ['account', A['zones'][0]['account'] + ' — ' + 'dd2832b36f171b815f84c8487aada36b'],
       ['zones_total_count (GET /zones result_info.total_count)', A['total_count']], ['zones_rows_read', len(A['zones'])], ['zones_read_at_utc', A['read_at']],
       ['traffic_source', 'Cloudflare GraphQL Analytics API — viewer.zones.httpRequests1dGroups (requests, cachedRequests, bytes, cachedBytes, threats, pageViews, encryptedRequests, uniques, responseStatusMap → 2xx/3xx/4xx/5xx, countryMap top 5 over the window)'],
       ['window', f'prior 30 days, daily, UTC: {days[0]} .. {days[-1]} inclusive (date_geq {T["window"][0]}, date_lt {T["window"][1]})'], ['traffic_read_at_utc', T['read_at']],
       ['token', 'read-only Analytics token, C:\\ALLOOLOO\\AGENT KEYS\\cloudflare-analytics.txt, read in-process, never printed'],
       ['joins', 'roster v1.1 (ALLOOLOO_TRAFFIC_zone-roster_v1.1_2026-10-24.xlsx: ZONES 308 + NEW-2026-10-24 112 rows, 109 include=Y, 3 TAKEN) · Aug 25 inventory (C:\\GREENCORE\\AUDITS\\CLOUDFLARE_DOMAINS\\CF-Zone-Inventory-2026-08-25.xlsx, 1,405 zones, family column) · Domains Master (C:\\ALLOOLOO\\COMPANY RECORDS\\Allooloo_Domains_Master.xlsx, Build sheet) · the estate-pages Worker custom-domain list (account read 2026-10-27, 190 zones)'],
       ['estate rule', 'roster v1.1 include=Y → allooloo · plan §5 shelf list (23), issuerwire.*, data-*/hunter-/hedge-/ark-agent → allooloo NEW-2026-10-28 · <cc>-issuer.ai, <cc>-eco.ai (Domains Master Allooloo), agentic-x402.*, zones served by the estate-pages Worker → allooloo, roster gap · agent-kg/familyofficeknowledgegraph/fokg/fo-kg/agent-padel/gsc-padel → office · agentic-zen/agenticzen → exclude · else gsc, family from the Aug 25 inventory, or by pattern (NG-9 §S3 grouping + the post-Aug-25 families), flagged'],
       ['lens rule (plan §6)', 'Company = allooloo.*, cm-kg.*, capitalmarketsknowledgegraph.* · KYP = getissuer.*/get-issuer.*, node door/twin, governance zone, product / standard, agentic-ask/-issuers/-coverage/-esg/-disclosure/-registries, issuer page/twin · Issuerwire = issuerwire.*/issuer-wire.* · Shelf = data-*, hunter-/hedge-/ark-agent · ECO = eco page/twin · Proof = agentic-radar.*, agentic-x402.* · Excluded = agentic-zen/agenticzen · unassigned = agentic-trades.*, agent-swap.* (§6 names no group)'],
       ['estate number', 'ESTATE!B2 = sum of the six Allooloo lens groups (Company, KYP, Issuerwire, Shelf, ECO, Proof), 30-day requests; unassigned zones shown beside it, not inside it'],
       ['failed reads', f'{len(T["failed"])} zones failed the GraphQL read; {len(T["missing_in_response"])} absent from the response — see FAILED'],
       ['written', datetime.datetime.utcnow().strftime('%Y-%m-%dT%H:%M:%SZ') + ' by CC (Claude Code), LENS-01']]
ws = wb.create_sheet('CONFIG'); [ws.append(r) for r in cfg]
for c in ws['A']: c.font = BF
for r in ws.iter_rows(min_row=1): r[1].font = F; r[1].alignment = Alignment(wrap_text=True, vertical='top')
ws.column_dimensions['A'].width = 44; ws.column_dimensions['B'].width = 150
# ---- ZONES (every zone, classified)
zh = ['zone', 'zone_id', 'estate', 'family', 'lens', 'include', 'flag', 'source', 'roster_v1.1', 'roster_family', 'inventory_2026-08-25_family', 'domains_master_group', 'status', 'plan', 'created_on', 'requests_30d', 'read']
zr = []
for z in sorted(Z, key=lambda z: (z['estate'], z['family'], z['zone'])):
    t = zone_tot.get(z['zone_id']); read = 'ok' if z['zone_id'] in T['daily'] else ('FAILED' if z['zone_id'] in T['failed'] else 'absent')
    zr.append([z['zone'], z['zone_id'], z['estate'], z['family'], z['lens'], z['include'], z['flag'], z['source'], z['roster_v11'], z['roster_family'], z['inventory_0825_family'], z['domains_master_group'], z['status'], z['plan'], z['created_on'], t['requests'] if t else None, read])
sheet(wb, 'ZONES', zh, zr, {'zone': 34, 'zone_id': 34, 'family': 26, 'lens': 24, 'flag': 30, 'source': 60, 'roster_v1.1': 22, 'roster_family': 22, 'inventory_2026-08-25_family': 22, 'domains_master_group': 30})
# ---- DAILY
dh = ['zone', 'date', 'estate', 'family', 'lens', 'requests', 'cached_requests', 'bytes', 'cached_bytes', 'threats', 'page_views', 'encrypted_requests', 'uniques', '2xx', '3xx', '4xx', '5xx', 'other_status']
sheet(wb, 'DAILY', dh, daily_rows, {'zone': 34, 'date': 12})
# ---- BY-ZONE (values from the read; 30-day)
bh = ['zone', 'estate', 'family', 'lens', 'include', 'requests_30d', 'cached_requests_30d', 'bytes_30d', 'threats_30d', 'page_views_30d', 'uniques_daily_sum', 'days_with_traffic', '2xx', '3xx', '4xx', '5xx', 'top_countries_30d', 'read']
br = []
for z in sorted(Z, key=lambda z: -(zone_tot.get(z['zone_id'], {}).get('requests', 0))):
    t = zone_tot.get(z['zone_id'], defaultdict(int)); read = 'ok' if z['zone_id'] in T['daily'] else ('FAILED' if z['zone_id'] in T['failed'] else 'absent')
    br.append([z['zone'], z['estate'], z['family'], z['lens'], z['include'], t['requests'], t['cachedRequests'], t['bytes'], t['threats'], t['pageViews'], t['uniques_daily_sum'], t['days_with_traffic'], t['2xx'], t['3xx'], t['4xx'], t['5xx'], top_countries(z['zone_id']), read])
bz = sheet(wb, 'BY-ZONE', bh, br, {'zone': 34, 'family': 26, 'lens': 24, 'top_countries_30d': 60})
n = len(br) + 1; RQ = f"'BY-ZONE'!$F$2:$F${n}"; EST = f"'BY-ZONE'!$B$2:$B${n}"; FAM = f"'BY-ZONE'!$C$2:$C${n}"; LENS = f"'BY-ZONE'!$D$2:$D${n}"; BY = f"'BY-ZONE'!$H$2:$H${n}"; TH = f"'BY-ZONE'!$I$2:$I${n}"
# ---- BY-FAMILY (formulas over BY-ZONE)
fams = OrderedDict()
for z in sorted(Z, key=lambda z: (z['estate'], z['family'])): fams.setdefault((z['estate'], z['family']), 0); fams[(z['estate'], z['family'])] += 1
fr = []
for i, ((e, f), c) in enumerate(fams.items(), 2):
    fr.append([e, f, c, f'=SUMIFS({RQ},{EST},A{i},{FAM},B{i})', f'=SUMIFS({BY},{EST},A{i},{FAM},B{i})', f'=SUMIFS({TH},{EST},A{i},{FAM},B{i})', f'=IF($D${len(fams) + 2}=0,0,D{i}/$D${len(fams) + 2})'])
fr.append(['TOTAL', '', f'=SUM(C2:C{len(fams) + 1})', f'=SUM(D2:D{len(fams) + 1})', f'=SUM(E2:E{len(fams) + 1})', f'=SUM(F2:F{len(fams) + 1})', ''])
wf = sheet(wb, 'BY-FAMILY', ['estate', 'family', 'zones', 'requests_30d', 'bytes_30d', 'threats_30d', 'share_of_account'], fr, {'family': 44})
for r in wf.iter_rows(min_row=2): r[6].number_format = '0.00%'; r[3].number_format = '#,##0'; r[4].number_format = '#,##0'
# ---- BY-LENS
lens_order = ['Company', 'KYP', 'Issuerwire', 'Shelf', 'ECO', 'Proof', 'unassigned (agentic-trades: §6 names no group)', 'unassigned (agent-swap: §6 names no group)', 'Excluded', '— (office)', '— (gsc)']
lens_counts = defaultdict(int)
for z in Z: lens_counts[z['lens']] += 1
lens_order = [l for l in lens_order if l in lens_counts] + [l for l in lens_counts if l not in lens_order]
lr = []
for i, l in enumerate(lens_order, 2):
    lr.append([l, 'allooloo' if l in ('Company', 'KYP', 'Issuerwire', 'Shelf', 'ECO', 'Proof') or l.startswith('unassigned') else ('exclude' if l == 'Excluded' else l.strip('— ()')), lens_counts[l], f'=SUMIFS({RQ},{LENS},A{i})', f'=SUMIFS({BY},{LENS},A{i})', f'=IF($D${len(lens_order) + 2}=0,0,D{i}/$D${len(lens_order) + 2})', 'in the estate number' if l in ('Company', 'KYP', 'Issuerwire', 'Shelf', 'ECO', 'Proof') else ('beside it — needs the CEO\'s word' if l.startswith('unassigned') else 'not counted')])
lr.append(['TOTAL (account)', '', f'=SUM(C2:C{len(lens_order) + 1})', f'=SUM(D2:D{len(lens_order) + 1})', f'=SUM(E2:E{len(lens_order) + 1})', '', ''])
wl = sheet(wb, 'BY-LENS', ['lens group', 'estate', 'zones', 'requests_30d', 'bytes_30d', 'share_of_account', 'counted'], lr, {'lens group': 48, 'counted': 34})
for r in wl.iter_rows(min_row=2): r[5].number_format = '0.00%'; r[3].number_format = '#,##0'; r[4].number_format = '#,##0'
# ---- BY-ESTATE
ests = ['allooloo', 'gsc', 'office', 'exclude']; er = []
for i, e in enumerate(ests, 2):
    er.append([e, f'=COUNTIFS({EST},A{i})', f'=SUMIFS({RQ},{EST},A{i})', f'=SUMIFS({BY},{EST},A{i})', f'=IF($C$6=0,0,C{i}/$C$6)'])
er.append(['TOTAL (account)', '=SUM(B2:B5)', '=SUM(C2:C5)', '=SUM(D2:D5)', ''])
we = sheet(wb, 'BY-ESTATE', ['estate', 'zones', 'requests_30d', 'bytes_30d', 'share_of_account'], er)
for r in we.iter_rows(min_row=2): r[4].number_format = '0.00%'; r[2].number_format = '#,##0'; r[3].number_format = '#,##0'
# ---- ESTATE (one number)
six = ['Company', 'KYP', 'Issuerwire', 'Shelf', 'ECO', 'Proof']
sr = [['ESTATE (Allooloo, six lens groups, 30-day requests)', '=' + '+'.join(f"SUMIFS({RQ},{LENS},\"{l}\")" for l in six), f'{days[0]} .. {days[-1]} UTC · the figure for the deck, the tear sheet, the RADAR tile and the digest line'],
      ['Allooloo share of the account', "=IF('BY-ESTATE'!C6=0,0,B2/'BY-ESTATE'!C6)", 'ESTATE ÷ all 2,134 zones'],
      ['Allooloo, all include=Y zones (adds the unassigned groups)', "=SUMIFS(" + RQ + "," + EST + ",\"allooloo\")", 'agentic-trades.* and agent-swap.* are outside the six groups until the CEO places them'],
      ['Account, every zone', "='BY-ESTATE'!C6", 'GSC + Allooloo + office + excluded'],
      ['Zones in the account', A['total_count'], 'GET /zones total_count'], ['Zones read with data', len(T['daily']), 'GraphQL response rows'], ['Zones whose read failed', len(T['failed']), 'see FAILED']]
we2 = sheet(wb, 'ESTATE', ['measure', 'value', 'note'], sr, {'measure': 60, 'value': 22, 'note': 100}, freeze=None)
for r in we2.iter_rows(min_row=2): r[1].number_format = '#,##0'; r[1].font = BF
we2['B3'].number_format = '0.00%'
# ---- FAILED
fr2 = [[by_id[t]['zone'], t, by_id[t]['estate'], e] for t, e in T['failed'].items()] + [[by_id[t]['zone'], t, by_id[t]['estate'], 'absent from the GraphQL response (zone returned no rows)'] for t in T['missing_in_response']]
sheet(wb, 'FAILED', ['zone', 'zone_id', 'estate', 'error'], fr2 or [['(none)', '', '', '']], {'zone': 34, 'zone_id': 34, 'error': 90})
wb.save(f'{OUT}/ZONES-{DATE}.xlsx'); print('wrote', f'{OUT}/ZONES-{DATE}.xlsx', 'daily rows', len(daily_rows))

# ---- roster v1.2
wb2 = openpyxl.Workbook(); wb2.remove(wb2.active)
cfg2 = [[r[0], r[1]] for r in R11['CONFIG']]
for r in cfg2:
    if r[0] == 'input': r[1] = 'this workbook — ZONES sheet, every zone in the account with estate · family · lens; the Allooloo readout = estate = allooloo, include = Y, lens in the six groups'
    if r[0] == 'status': r[1] = f'roster v1.2 · {DATE} · {len(Z)} zones in the account, {sum(1 for z in Z if z["estate"]=="allooloo")} Allooloo · written by LENS-01 from roster v1.1 + the Aug 25 inventory + the account read; nothing on the wire changed'
cfg2 += [['lens groups (plan §6)', 'Company · KYP · Issuerwire · Shelf · ECO · Proof; Excluded; unassigned = agentic-trades.*, agent-swap.* (§6 names no group)'], ['prior', 'ALLOOLOO_TRAFFIC_zone-roster_v1.1_2026-10-24.xlsx (ZONES + NEW-2026-10-24); the 4 eu-eco zones were listed on both sheets there (417 counted, 413 distinct)']]
ws = wb2.create_sheet('CONFIG'); [ws.append(r) for r in cfg2]
for c in ws['A']: c.font = BF
for r in ws.iter_rows(min_row=1): r[1].font = F; r[1].alignment = Alignment(wrap_text=True, vertical='top')
ws.column_dimensions['A'].width = 30; ws.column_dimensions['B'].width = 150
rh = ['zone', 'estate', 'family', 'lens', 'include', 'tld', 'cloudflare_account', 'zone_id', 'created_on', 'roster_v1.1', 'flag', 'source', 'requests_30d_2026-10-28']
rr = [[z['zone'], z['estate'], z['family'], z['lens'], z['include'], z['zone'].split('.', 1)[1], 'GSC and Allooloo AI Agent Stack (dd2832b3…)', z['zone_id'], z['created_on'], z['roster_v11'], z['flag'], z['source'], zone_tot.get(z['zone_id'], {}).get('requests')]
      for z in sorted(Z, key=lambda z: ({'allooloo': 0, 'office': 1, 'exclude': 2, 'gsc': 3}[z['estate']], z['lens'], z['family'], z['zone']))]
sheet(wb2, 'ZONES', rh, rr, {'zone': 34, 'family': 26, 'lens': 24, 'cloudflare_account': 36, 'zone_id': 34, 'roster_v1.1': 22, 'flag': 30, 'source': 60})
newr = [[z['zone'], z['family'], z['lens'], z['flag'], z['source'], z['created_on']] for z in sorted(Z, key=lambda z: z['zone']) if z['estate'] == 'allooloo' and z['flag']]
sheet(wb2, f'NEW-{DATE}', ['zone', 'family', 'lens', 'flag', 'source', 'created_on'], newr, {'zone': 34, 'family': 26, 'lens': 24, 'flag': 30, 'source': 70})
famr = []
afams = OrderedDict()
for z in sorted(Z, key=lambda z: z['family']):
    if z['estate'] == 'allooloo': afams.setdefault(z['family'], [0, 0, set()]); afams[z['family']][0] += 1; afams[z['family']][1] += 1 if z['flag'] else 0; afams[z['family']][2].add(z['lens'])
for i, (f, (c, nw, ls)) in enumerate(afams.items(), 2): famr.append([f, c, nw, ' · '.join(sorted(ls)), f"=COUNTIFS(ZONES!$C$2:$C${len(rr)+1},A{i},ZONES!$B$2:$B${len(rr)+1},\"allooloo\")"])
famr.append(['TOTAL (Allooloo)', f'=SUM(B2:B{len(afams)+1})', f'=SUM(C2:C{len(afams)+1})', '', f'=SUM(E2:E{len(afams)+1})'])
sheet(wb2, 'FAMILIES', ['family', 'zones', f'of which new {DATE}', 'lens', 'check (COUNTIFS on ZONES)'], famr, {'family': 30, 'lens': 40, 'check (COUNTIFS on ZONES)': 26})
lensr = [[l, sum(1 for z in Z if z['lens'] == l), f"=COUNTIFS(ZONES!$D$2:$D${len(rr)+1},A{i})"] for i, l in enumerate(lens_order, 2)]
sheet(wb2, 'LENS', ['lens group', 'zones', 'check (COUNTIFS on ZONES)'], lensr, {'lens group': 48, 'check (COUNTIFS on ZONES)': 26})
wb2.save(f'{OUT}/ALLOOLOO_TRAFFIC_zone-roster_v1.2_{DATE}.xlsx'); print('wrote roster v1.2', len(rr), 'rows; new', len(newr))
